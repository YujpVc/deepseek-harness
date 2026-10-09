"""Keyless native-chart CLI, data-integrity and failure-path checks."""
import copy
import json
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from pptx import Presentation
from pptx.util import Inches

from build_deck import build, build_probation_deck_v2
from charts import KINDS, resolve_chart

ROOT = Path(__file__).resolve().parent


def presentation_snapshot(prs):
    """Record native data, chart kinds and visible text, excluding generated ids."""
    result = []
    for slide in prs.slides:
        charts = []
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            chart = shape.chart
            data = []
            for series in chart.series:
                item = {'name': series.name, 'values': list(series.values)}
                for tag, field in [('xVal', 'x'), ('bubbleSize', 'size')]:
                    nodes = series._element.xpath(f'c:{tag}/c:numRef/c:numCache/c:pt/c:v')
                    if nodes:
                        item[field] = [float(node.text) for node in nodes]
                data.append(item)
            charts.append({'type': int(chart.chart_type), 'series': data})
        result.append({'charts': charts, 'text': [s.text for s in slide.shapes
                                                 if s.has_text_frame and s.text.strip()]})
    return result


class NativeChartTest(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.base = Path(tmp.name)
        self.output = self.base / 'deck.pptx'
        self.chart = {'kind': 'column', 'categories': ['A', 'B', 'C'],
                      'series': [{'name': 'Result', 'values': [2, None, 4]}]}

    def spec(self, chart):
        return {'meta': {'title': 'Native data', 'theme': 'paper'},
                'slides': [{'layout': 'chart', 'title': 'Result', 'chart': chart}]}

    def test_all_kinds_roundtrip_with_embedded_workbooks(self):
        for kind in KINDS:
            with self.subTest(kind=kind):
                chart = copy.deepcopy(self.chart)
                chart['kind'] = kind
                chart['series'][0]['values'] = [2, 3, 4]
                if kind in {'scatter', 'bubble'}:
                    chart = {'kind': kind, 'series': [{'name': 'Result', 'points':
                             [[1, 2], [3, 4]] if kind == 'scatter' else [[1, 2, 5], [3, 4, 9]]}]}
                build(self.spec(chart), self.base, self.output)
                prs = Presentation(self.output)
                native = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
                self.assertEqual(native.chart_type, KINDS[kind])
                self.assertEqual(list(native.series[0].values), [2, 4] if kind in {'scatter', 'bubble'} else [2, 3, 4])
                with zipfile.ZipFile(self.output) as archive:
                    workbooks = [n for n in archive.namelist() if n.endswith('.xlsx')]
                    self.assertEqual(len(workbooks), 1)
                    with zipfile.ZipFile(io.BytesIO(archive.read(workbooks[0]))) as workbook:
                        self.assertIn('xl/worksheets/sheet1.xml', workbook.namelist())
                if kind.startswith('percent_'):
                    self.assertFalse(native.plots[0].has_data_labels)

    def test_missing_measurements_remain_gaps(self):
        self.chart['kind'] = 'line'
        build(self.spec(self.chart), self.base, self.output)
        native = next(s.chart for s in Presentation(self.output).slides[0].shapes if s.has_chart)
        self.assertEqual(list(native.series[0].values), [2, None, 4])
        self.assertEqual(native._chartSpace.xpath('c:chart/c:dispBlanksAs')[0].get('val'), 'gap')

    def test_csv_bom_sorting_and_gaps_preserve_pairing(self):
        (self.base / 'source.csv').write_text('label,value\nA,2\nB,\nC,5\n', encoding='utf-8-sig')
        chart = {'kind': 'bar', 'data_file': 'source.csv', 'category_column': 'label',
                 'series': [{'name': 'Value', 'column': 'value'}], 'sort': 'descending'}
        original = copy.deepcopy(chart)
        resolved = resolve_chart(chart, self.base)
        self.assertEqual(resolved['categories'], ['C', 'A', 'B'])
        self.assertEqual(resolved['series'][0]['values'], [5, 2, None])
        self.assertEqual(chart, original)
        build(self.spec(chart), self.base, self.output)
        native = next(s.chart for s in Presentation(self.output).slides[0].shapes if s.has_chart)
        self.assertEqual([c.label for c in native.plots[0].categories], ['C', 'A', 'B'])
        self.assertEqual(list(native.series[0].values), [5, 2, None])

    def test_invalid_inputs_do_not_replace_existing_output(self):
        invalid = [
            {'kind': 'unknown'}, {'kind': []}, {'series': [None]},
            {'series': [{'name': 'x', 'values': [True, 2, 3]}]},
            {'series': [{'name': 'x', 'values': [1, float('nan'), 3]}]},
            {'series': [{'name': 'x', 'values': [None, None, None]}]},
            {'series': [{'name': 'x', 'values': [1]}]},
            {'kind': 'pie', 'series': [{'name': 'x', 'values': [1, -2, 3]}]},
            {'kind': 'percent_column', 'series': [{'name': 'x', 'values': [0, 0, 0]}]},
            {'kind': 'percent_bar', 'number_format': '0%'},
            {'kind': 'percent_bar', 'y_axis': {'max': 100}},
            {'y_axis': {'min': 3}}, {'y_axis': {'min': 2, 'max': 1}},
            {'y_axis': {'major_unit': 0}}, {'x_axis': {'min': 0}},
            {'x_axis': {'font_size': 3}}, {'highlight': ['missing']},
            {'series': [{'name': 'x', 'values': [1, 2, 3], 'color': '#abc'}]},
            {'kind': 'bubble', 'series': [{'name': 'x', 'points': [[1, 2, -3]]}]},
            {'data_labels': 'yes'}, {'data_file': None},
        ]
        self.output.write_bytes(b'keep existing deck')
        for patch in invalid:
            with self.subTest(patch=patch):
                chart = copy.deepcopy(self.chart)
                chart.update(patch)
                with self.assertRaises((ValueError, OSError)):
                    build(self.spec(chart), self.base, self.output)
                self.assertEqual(self.output.read_bytes(), b'keep existing deck')

    def test_stacked_axis_cannot_hide_totals(self):
        chart = {'kind': 'stacked_column', 'categories': ['A'],
                 'series': [{'name': 'one', 'values': [4]}, {'name': 'two', 'values': [5]}],
                 'y_axis': {'max': 8}}
        with self.assertRaisesRegex(ValueError, 'hide source'):
            resolve_chart(chart, self.base)

    def test_csv_malformed_headers_rows_and_mixed_sources(self):
        chart = {'data_file': 'bad.csv', 'category_column': 'label',
                 'series': [{'name': 'x', 'column': 'value'}]}
        for content in ('label,value\nA\n', 'label,value\nA,2,extra\n',
                        'label,value,value\nA,2,3\n', 'label,value\nA,wrong\n'):
            (self.base / 'bad.csv').write_text(content)
            with self.subTest(content=content), self.assertRaises(ValueError):
                resolve_chart(chart, self.base)
        chart['categories'] = ['A']
        with self.assertRaisesRegex(ValueError, 'mixed'):
            resolve_chart(chart, self.base)

    def test_template_expansion_uses_the_same_native_charts(self):
        template = self.base / 'reference.pptx'
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
        prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = 'Keep cover'
        cover_xml = prs.slides[0]._element.xml
        prs.save(template)
        spec = self.spec(self.chart)
        spec['slides'].insert(0, {'type': 'cover'})
        build_probation_deck_v2(spec, self.base, template, self.output)
        output = Presentation(self.output)
        self.assertEqual(output.slides[0]._element.xml, cover_xml)
        self.assertEqual(sum(s.has_chart for s in output.slides[1].shapes), 1)

    def test_dashboard_capacity_and_dense_labels_fail_without_overwrite(self):
        spec = self.spec(self.chart)
        spec['slides'][0].update(layout='chart_dashboard', charts=[self.chart])
        with self.assertRaises(ValueError):
            build(spec, self.base, self.output)
        dense = {'kind': 'bar', 'categories': [str(i) for i in range(30)],
                 'series': [{'name': 'x', 'values': list(range(30))}]}
        self.output.write_bytes(b'keep')
        with self.assertRaisesRegex(ValueError, 'panel height'):
            build(self.spec(dense), self.base, self.output)
        self.assertEqual(self.output.read_bytes(), b'keep')

    def test_three_chart_dashboard_and_axes_remain_editable(self):
        chart = copy.deepcopy(self.chart)
        chart['y_axis'] = {'title': 'Time (s)', 'min': 0, 'max': 5}
        spec = self.spec(chart)
        spec['slides'][0].update(layout='chart_dashboard', charts=[chart, chart, chart])
        build(spec, self.base, self.output)
        shapes = [s for s in Presentation(self.output).slides[0].shapes if s.has_chart]
        self.assertEqual(len(shapes), 3)
        self.assertGreater(shapes[2].width, shapes[0].width * 1.9)
        self.assertEqual(shapes[0].chart.value_axis.axis_title.text_frame.text, 'Time (s)')
        self.assertEqual(shapes[0].chart.value_axis.maximum_scale, 5)

    def test_dark_theme_sets_chart_background(self):
        spec = self.spec(self.chart)
        spec['meta']['theme'] = 'midnight'
        build(spec, self.base, self.output)
        chart = next(s.chart for s in Presentation(self.output).slides[0].shapes if s.has_chart)
        self.assertEqual(chart._chartSpace.xpath('c:spPr/a:solidFill/a:srgbClr')[0].get('val'), '17243B')

    def test_long_trend_uses_sparse_ticks_without_dropping_values(self):
        chart = {'kind': 'line', 'categories': [str(i) for i in range(40)],
                 'series': [{'name': 'x', 'values': list(range(40))}]}
        build(self.spec(chart), self.base, self.output)
        native = next(s.chart for s in Presentation(self.output).slides[0].shapes if s.has_chart)
        self.assertEqual(list(native.series[0].values), list(range(40)))
        self.assertGreater(int(native.category_axis._element.xpath('c:tickLblSkip')[0].get('val')), 1)

    def test_example_cli_matches_reviewed_snapshot(self):
        result = subprocess.run([sys.executable, str(ROOT / 'build_deck.py'),
                                 str(ROOT / 'examples/chart_report.json'), str(self.output)],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        expected = json.loads((ROOT / 'examples/chart_report.expected.json').read_text())
        self.assertEqual(presentation_snapshot(Presentation(self.output)), expected)
        gate = subprocess.run([sys.executable, str(ROOT / 'quality_check.py'), str(self.output)],
                              capture_output=True, text=True)
        self.assertEqual(gate.returncode, 0, gate.stdout + gate.stderr)


if __name__ == '__main__':
    unittest.main()
