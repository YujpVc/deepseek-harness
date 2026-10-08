#!/usr/bin/env python3
"""Fail-fast structural quality gate for generated PowerPoint decks.

This is deliberately independent of LibreOffice: it catches the common failure
where an agent delivers a text-only ``pptx_create`` draft before visual review.
It does not replace rendering; it prevents an obviously incomplete deck from
being presented as finished when rendering is unavailable.
"""
from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def inspect(path: Path) -> dict:
    prs = Presentation(str(path))
    slides = []
    for index, slide in enumerate(prs.slides, 1):
        images = 0
        tables = 0
        charts = 0
        text_boxes = 0
        auto_shapes = 0
        diagram_shapes = 0
        for shape in slide.shapes:
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                images += 1
            elif getattr(shape, "has_table", False):
                tables += 1
            elif getattr(shape, "has_chart", False):
                charts += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                auto_shapes += 1
                area_ratio = shape.width * shape.height / (prs.slide_width * prs.slide_height)
                if 0.01 <= area_ratio <= 0.75:
                    diagram_shapes += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_boxes += 1
        slides.append(
            {
                "slide": index,
                "images": images,
                "tables": tables,
                "charts": charts,
                "text_boxes": text_boxes,
                "auto_shapes": auto_shapes,
                "diagram_shapes": diagram_shapes,
                "shapes": len(slide.shapes),
            }
        )

    width = prs.slide_width / 914400
    height = prs.slide_height / 914400
    visual_slides = sum(
        1 for item in slides
        if item["images"] or item["tables"] or item["charts"] or item["diagram_shapes"] >= 3
    )
    signatures = {
        (
            item["images"],
            item["tables"],
            item["charts"],
            min(item["text_boxes"], 14),
            min(item["auto_shapes"], 8),
            min(item["shapes"], 16),
        )
        for item in slides
    }
    return {
        "file": str(path),
        "slides": len(slides),
        "width_in": round(width, 3),
        "height_in": round(height, 3),
        "aspect_ratio": round(width / height, 4) if height else None,
        "visual_slides": visual_slides,
        "layout_signatures": len(signatures),
        "slide_stats": slides,
    }


def failures(report: dict) -> list[str]:
    problems: list[str] = []
    if report["slides"] >= 5:
        if abs(report["aspect_ratio"] - 16 / 9) > 0.03:
            problems.append(
                f"页面比例为 {report['width_in']}×{report['height_in']}，不是 16:9"
            )
        required_visual = max(3, math.ceil(report["slides"] * 0.4))
        if report["visual_slides"] < required_visual:
            problems.append(
                f"仅 {report['visual_slides']}/{report['slides']} 页含图片、表格、图表或原生图示，至少需要 {required_visual} 页"
            )
        if report["layout_signatures"] < 4:
            problems.append(
                f"仅检测到 {report['layout_signatures']} 种页面结构，至少需要 4 种"
            )
        text_only = [
            item["slide"]
            for item in report["slide_stats"]
            if not (item["images"] or item["tables"] or item["charts"])
            and item["text_boxes"] <= 3
            and item["diagram_shapes"] == 0
        ]
        if len(text_only) > max(1, report["slides"] // 4):
            problems.append(
                "以下页面几乎只有文字框，疑似未完成排版：" + ", ".join(map(str, text_only))
            )
    return problems


def compare_template(report: dict, template_path: Path) -> list[str]:
    """Require each output page to retain its source template's structure."""
    baseline = inspect(template_path)
    problems: list[str] = []
    if report["slides"] != baseline["slides"]:
        problems.append(
            f"模板页数为 {baseline['slides']}，输出为 {report['slides']}，模板页未完整继承"
        )
        return problems
    source = Presentation(str(template_path))
    result = Presentation(report["file"])
    for index, (before, after, source_slide, result_slide) in enumerate(
        zip(baseline["slide_stats"], report["slide_stats"], source.slides, result.slides), 1
    ):
        source_types = [int(shape.shape_type) for shape in source_slide.shapes]
        result_types = [int(shape.shape_type) for shape in result_slide.shapes]
        if source_types != result_types:
            problems.append(f"第 {index} 页形状结构与模板不一致")
        if source_slide.slide_layout.name != result_slide.slide_layout.name:
            problems.append(f"第 {index} 页布局与模板不一致")
        if before["images"] != after["images"]:
            problems.append(f"第 {index} 页图片数量与模板不一致")
    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--template", type=Path,
                        help="compare page count, layout, shapes and image count against a source template")
    args = parser.parse_args()
    try:
        report = inspect(args.pptx)
    except Exception as exc:  # clear failure for the agent, without a traceback wall
        print(f"quality_check: 无法读取 {args.pptx}: {exc}", file=sys.stderr)
        return 2
    problems = failures(report)
    if args.template:
        if not args.template.is_file():
            print(f"quality_check: template not found: {args.template}", file=sys.stderr)
            return 2
        template_problems = compare_template(report, args.template)
        report["template"] = str(args.template.resolve())
        report["template_failures"] = template_problems
        problems.extend(template_problems)
    if args.json:
        report["failures"] = problems
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(
            f"页数={report['slides']} 比例={report['width_in']}×{report['height_in']} "
            f"视觉页={report['visual_slides']} 版式签名={report['layout_signatures']}"
        )
        if problems:
            print("FAIL")
            for problem in problems:
                print(f"- {problem}")
        else:
            print("PASS")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
