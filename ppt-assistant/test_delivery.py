"""Keyless tests for notes preservation, animation, vector formulas and delivery."""
import json
import shutil
import subprocess
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from delivery import audit, sha256, sync_files
from formula import add_formula
from speaker_notes import embed, export_script, parse_script, resolve_part
from video_gif import choose_speed, convert, gif_info, insert


class DeliveryTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.deck = self.root / "original.pptx"
        self.output = self.root / "output.pptx"
        self.script = self.root / "script.md"
        self.gif = self.root / "demo.gif"
        frames = [Image.new("RGB", (80, 120), color) for color in ("red", "blue", "green")]
        frames[0].save(self.gif, save_all=True, append_images=frames[1:], duration=100, loop=0)
        prs = Presentation()
        first = prs.slides.add_slide(prs.slide_layouts[6])
        first.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Evidence"
        first.shapes.add_picture(str(self.gif), Inches(1), Inches(2), Inches(2), Inches(2)).name = "Demo"
        second = prs.slides.add_slide(prs.slide_layouts[6])
        second.notes_slide.notes_text_frame.text = "Old notes"
        prs.save(self.deck)
        self.script.write_text("# Script\n\n## P1｜Evidence（20 秒）\n\nTested results.\n\n"
                               "## P2｜Plan (15 seconds)\n\nNext step.\n", encoding="utf-8")

    def test_notes_only_update_preserves_all_content_and_reads_back(self):
        report = embed(self.deck, self.script, self.output)
        self.assertEqual(report["budget_seconds"], 35)
        result = audit(self.output, self.deck, self.script, (1,))
        self.assertTrue(result["content_preserved"])
        self.assertTrue(result["notes_match_markdown"])
        self.assertEqual(result["gif_pages"]["1"][0]["frames"], 3)
        first = sha256(self.output)
        embed(self.output, self.script, self.output)
        self.assertEqual(sha256(self.output), first)
        with ZipFile(self.deck) as before, ZipFile(self.output) as after:
            for name in before.namelist():
                if name.startswith(("ppt/slides/slide", "ppt/media/", "ppt/slideMasters/")):
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_deck_without_any_notes_or_master_gets_valid_notes(self):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(self.deck)
        self.script.write_text("## P1｜Opening\n\nHello.\n", encoding="utf-8")
        embed(self.deck, self.script, self.output)
        self.assertEqual(Presentation(self.output).slides[0].notes_slide.notes_text_frame.text,
                         "P1｜Opening\n\nHello.")
        self.assertTrue(audit(self.output, self.deck, self.script)["content_preserved"])

    def test_slide_order_is_resolved_from_presentation_relationships(self):
        prs = Presentation(self.deck)
        ids = prs.slides._sldIdLst
        ids.insert(0, ids[-1])
        prs.save(self.deck)
        embed(self.deck, self.script, self.output)
        self.assertTrue(Presentation(self.output).slides[0].notes_slide.notes_text_frame.text.startswith("P1｜Evidence"))
        self.assertTrue(audit(self.output, self.deck)["content_preserved"])

    def test_invalid_script_cannot_overwrite_existing_output(self):
        self.output.write_bytes(b"keep this")
        invalid = ["## P1｜Only\n\nOne.\n", "## P1｜Empty\n\n## P2｜Second\n\nTwo.\n",
                   "## P1｜Draft\n\n[TODO speaker script]\n\n## P2｜Two\n\nTwo.\n",
                   "## P1｜One\n\nOne.\n\n## P1｜Duplicate\n\nTwo.\n"]
        for content in invalid:
            self.script.write_text(content)
            with self.assertRaises(ValueError):
                embed(self.deck, self.script, self.output)
            self.assertEqual(self.output.read_bytes(), b"keep this")

    def test_export_does_not_invent_script_and_marks_missing_notes(self):
        export_script(self.deck, self.script)
        text = self.script.read_text()
        self.assertIn("Evidence", text)
        self.assertIn("[TODO speaker script]", text)
        self.assertIn("Old notes", text)
        with self.assertRaises(ValueError):
            export_script(self.deck, self.deck)

    def test_comment_headings_are_excluded_and_absolute_targets_resolve(self):
        self.script.write_text("<!--\n## P9｜Evidence only\nIgnored.\n-->\n"
                               "## P1｜Opening\nHello.\n## P2｜Plan\nNext.\n")
        self.assertEqual(set(parse_script(self.script, 2)[0]), {1, 2})
        self.assertEqual(resolve_part("ppt/slides/slide1.xml", "/ppt/notesSlides/notesSlide1.xml"),
                         "ppt/notesSlides/notesSlide1.xml")

    def test_audit_rejects_changed_text_and_notes(self):
        embed(self.deck, self.script, self.output)
        prs = Presentation(self.output)
        prs.slides[0].shapes[0].text = "Changed result"
        prs.save(self.output)
        with self.assertRaisesRegex(ValueError, "outside notes"):
            audit(self.output, self.deck)
        embed(self.deck, self.script, self.output)
        prs = Presentation(self.output)
        prs.slides[0].notes_slide.notes_text_frame.text = "Wrong script"
        prs.save(self.output)
        with self.assertRaisesRegex(ValueError, "differ from Markdown"):
            audit(self.output, script=self.script)

    def test_missing_or_static_gif_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "missing animated GIF"):
            audit(self.deck, gif_pages=(2,))
        stream = BytesIO()
        Image.new("RGB", (10, 10)).save(stream, format="GIF")
        with self.assertRaisesRegex(ValueError, "multi-frame"):
            gif_info(stream.getvalue())

    def test_unused_gif_relationship_does_not_count_as_page_animation(self):
        prs = Presentation(self.deck)
        shape = prs.slides[0].shapes[1]
        shape._element.getparent().remove(shape._element)
        prs.save(self.output)
        with ZipFile(self.output) as archive:
            self.assertTrue(any(name.endswith(".gif") for name in archive.namelist()))
        with self.assertRaisesRegex(ValueError, "missing animated GIF"):
            audit(self.output, gif_pages=(1,))

    def test_insert_retains_animation_and_adds_editable_speed_marker(self):
        insert(self.deck, self.gif, self.output, 1, "Demo", 10)
        slide = Presentation(self.output).slides[0]
        self.assertEqual(next(s for s in slide.shapes if s.name == "Demo").image.blob, self.gif.read_bytes())
        self.assertEqual(next(s for s in slide.shapes if s.name == "Demo speed").text, "10X")
        self.assertEqual(audit(self.output, gif_pages=(1,))["gif_pages"]["1"][0]["frames"], 3)

    def test_speed_selection_uses_integer_multiples_and_handles_short_video(self):
        self.assertEqual(choose_speed(108), 10)
        self.assertEqual(choose_speed(196), 20)
        self.assertEqual(choose_speed(30), 3)
        self.assertEqual(choose_speed(4), 1)
        for duration in (0, float("inf"), float("nan")):
            with self.assertRaises(ValueError):
                choose_speed(duration)

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "requires ffmpeg/ffprobe")
    def test_conversion_keeps_source_ending_instead_of_trimming(self):
        source = self.root / "source.mp4"
        subprocess.run(["ffmpeg", "-v", "error", "-f", "lavfi", "-i", "color=red:s=80x120:d=1:r=20",
                        "-f", "lavfi", "-i", "color=blue:s=80x120:d=1:r=20", "-filter_complex",
                        "[0:v][1:v]concat=n=2:v=1:a=0[v]", "-map", "[v]", str(source)], check=True)
        report = convert(source, self.gif, speed=2, fps=10, size=120)
        self.assertEqual(report["source_duration_seconds"], 2)
        self.assertAlmostEqual(report["duration_seconds"], 1, delta=.15)
        with Image.open(self.gif) as image:
            image.seek(image.n_frames - 1)
            r, g, b = image.convert("RGB").getpixel((40, 60))
            self.assertGreater(b, r + 100)

    def test_formula_is_native_geometry_with_curves_and_no_image(self):
        prs = Presentation(self.deck)
        shape = add_formula(prs.slides[0], r"$T_{\Delta}=T_t^{-1}T_{t+1}$",
                            Inches(1), Inches(4), Inches(6), Inches(1))
        self.assertIn("cubicBezTo", shape._element.xml)
        self.assertNotIn("blip", shape._element.xml)
        prs.save(self.output)
        self.assertEqual(Presentation(self.output).slides[0].shapes[-1].name, "Formula")
        with self.assertRaises(ValueError):
            add_formula(prs.slides[0], r"$\unsupported{x}$", 0, 0, 100, 100)

    def test_sync_checks_remote_sha_and_does_not_delete_or_store_passwords(self):
        calls = []
        destination = "/remote folder/a'b"
        def run(command, **kwargs):
            calls.append(command)
            stdout = f"{sha256(self.script)}  {destination}/script.md\n" if "sha256sum" in command[-1] else ""
            return subprocess.CompletedProcess(command, 0, stdout=stdout)
        with patch("delivery.subprocess.run", side_effect=run):
            self.assertTrue(sync_files(self.root, [self.script], "desktop", destination)["sha256_verified"])
        self.assertEqual(len(calls), 3)
        self.assertNotIn("--delete", calls[1])
        self.assertIn("--protect-args", calls[1])
        self.assertIn("'\"'\"'", calls[-1][-1])
        with patch("delivery.subprocess.run", return_value=subprocess.CompletedProcess([], 0, stdout="bad  file\n")):
            with self.assertRaisesRegex(ValueError, "SHA256 mismatch"):
                sync_files(self.root, [self.script], "desktop", destination)
        with self.assertRaises(ValueError):
            sync_files(self.root, [self.script], "-bad", destination)


if __name__ == "__main__":
    unittest.main()
