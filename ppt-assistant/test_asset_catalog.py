"""Keyless tests for deterministic asset inventories and editable slide reuse."""
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from asset_catalog import scan, verify, write_manifest
from slide_ops import copy_slides


class AssetCatalogTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_scan_manifest_is_stable_and_verifiable(self):
        (self.root / "nested").mkdir()
        (self.root / "nested/data.json").write_text('{"ok": true}\n', encoding="utf-8")
        Image.new("RGB", (12, 8), "navy").save(self.root / "cover.png")
        output = self.root / "manifest.json"
        markdown = self.root / "index.md"
        manifest = write_manifest(self.root, output, markdown)
        self.assertEqual([item["path"] for item in manifest["files"]], ["cover.png", "nested/data.json"])
        self.assertEqual(verify(self.root, output, strict=True)["verified"], 2)
        self.assertIn("cover.png", markdown.read_text(encoding="utf-8"))
        (self.root / "nested/data.json").write_text('{"changed": true}\n', encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "SHA256 mismatch"):
            verify(self.root, output)

    def test_scan_skips_hidden_and_symlink_files(self):
        (self.root / ".private").write_text("secret", encoding="utf-8")
        (self.root / "visible.txt").write_text("ok", encoding="utf-8")
        (self.root / "alias.txt").symlink_to(self.root / "visible.txt")
        self.assertEqual([item["path"] for item in scan(self.root)["files"]], ["visible.txt"])

    def test_strict_detects_additions_and_rejects_symlink_replacements(self):
        source = self.root / "source.txt"
        source.write_text("record", encoding="utf-8")
        manifest = self.root / "manifest.json"
        write_manifest(self.root, manifest)
        extra = self.root / "extra.txt"
        extra.write_text("record", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "strict inventory differs"):
            verify(self.root, manifest, strict=True)
        source.unlink()
        source.symlink_to(extra)
        with self.assertRaisesRegex(ValueError, "missing or leaves root"):
            verify(self.root, manifest)

    def test_manifest_rejects_traversal_and_noncanonical_names(self):
        manifest = self.root / "manifest.json"
        for name in ("../outside.txt", "/etc/passwd", "nested/../file", "./file", "", "nested//file"):
            with self.subTest(name=name):
                manifest.write_text(json.dumps({"version": 1, "files": [{"path": name}]}), encoding="utf-8")
                with self.assertRaises(ValueError):
                    verify(self.root, manifest)
        manifest.write_text(json.dumps({"version": 1, "files": [], "excluded": ["../outside"]}), encoding="utf-8")
        with self.assertRaises(ValueError):
            verify(self.root, manifest, strict=True)


class SlideOpsTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.pptx"
        self.target = self.root / "target.pptx"
        self.output = self.root / "merged.pptx"
        image = self.root / "media.png"
        Image.new("RGB", (20, 20), "orange").save(image)
        source = Presentation()
        slide = source.slides.add_slide(source.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Reusable page"
        slide.shapes.add_picture(str(image), Inches(1), Inches(2), Inches(1), Inches(1))
        slide.notes_slide.notes_text_frame.text = "Reusable speaker note"
        source.save(self.source)
        target = Presentation()
        target.slides.add_slide(target.slide_layouts[6]).shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)).text = "Target page"
        target.save(self.target)

    def test_copy_keeps_editable_content_media_and_notes(self):
        result = copy_slides(self.source, self.target, self.output, [1], position=0)
        self.assertEqual(result["output_pages"], 2)
        merged = Presentation(self.output)
        self.assertEqual(merged.slides[0].shapes[0].text, "Reusable page")
        self.assertTrue(any(shape.shape_type == 13 for shape in merged.slides[0].shapes))
        self.assertEqual(merged.slides[0].notes_slide.notes_text_frame.text, "Reusable speaker note")
        self.assertEqual(merged.slides[1].shapes[0].text, "Target page")

    def test_invalid_page_does_not_replace_existing_output(self):
        self.output.write_bytes(b"keep")
        with self.assertRaises(ValueError):
            copy_slides(self.source, self.target, self.output, [2])
        self.assertEqual(self.output.read_bytes(), b"keep")

    def test_colliding_images_charts_and_workbooks_retain_distinct_data(self):
        source, target = Presentation(self.source), Presentation(self.target)
        target_image = self.root / "target.png"
        Image.new("RGB", (20, 20), "blue").save(target_image)
        target.slides[0].shapes.add_picture(str(target_image), Inches(1), Inches(2), Inches(1), Inches(1))
        for prs, value in ((source, 9), (target, 3)):
            data = CategoryChartData()
            data.categories = ["Value"]
            data.add_series("Series", [value])
            prs.slides[0].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(2), Inches(2),
                                         Inches(3), Inches(2), data)
        animation = self.root / "demo.gif"
        Image.new("RGB", (20, 20), "red").save(animation, save_all=True,
            append_images=[Image.new("RGB", (20, 20), "green")], duration=100, loop=0)
        source.slides[0].shapes.add_picture(str(animation), Inches(6), Inches(2), Inches(1), Inches(1))
        source.save(self.source)
        target.save(self.target)
        copy_slides(self.source, self.target, self.output, [1, 1])
        with ZipFile(self.output) as archive:
            self.assertEqual(len(archive.namelist()), len(set(archive.namelist())))
        merged = Presentation(self.output)
        self.assertEqual(merged.slides[0].shapes[1].image.blob, target_image.read_bytes())
        self.assertEqual(merged.slides[1].shapes[1].image.blob, (self.root / "media.png").read_bytes())
        self.assertEqual(merged.slides[1].shapes[-1].image.blob, animation.read_bytes())
        for page, value in ((0, 3), (1, 9), (2, 9)):
            chart = next(shape.chart for shape in merged.slides[page].shapes if shape.has_chart)
            self.assertEqual(chart.series[0].values, (value,))
            original = target if page == 0 else source
            original_chart = next(shape.chart for shape in original.slides[0].shapes if shape.has_chart)
            self.assertEqual(chart.part.chart_workbook.xlsx_part.blob,
                             original_chart.part.chart_workbook.xlsx_part.blob)

    def test_dimension_mismatch_and_internal_slide_links_leave_output_untouched(self):
        self.output.write_bytes(b"keep")
        target = Presentation(self.target)
        target.slide_width += Inches(1)
        target.save(self.target)
        with self.assertRaisesRegex(ValueError, "dimensions must match"):
            copy_slides(self.source, self.target, self.output, [1])
        source = Presentation(self.source)
        source.slides[0].shapes[0].click_action.target_slide = source.slides[0]
        source.save(self.source)
        target.slide_width -= Inches(1)
        target.save(self.target)
        with self.assertRaisesRegex(ValueError, "internal slide links"):
            copy_slides(self.source, self.target, self.output, [1])
        self.assertEqual(self.output.read_bytes(), b"keep")


if __name__ == "__main__":
    unittest.main()
