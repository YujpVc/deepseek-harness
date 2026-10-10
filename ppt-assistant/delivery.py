#!/usr/bin/env python3
"""Audit a notes-only deck and synchronize explicit files with remote SHA256 checks."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shlex
import subprocess
import tempfile
from pathlib import Path, PurePosixPath
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from speaker_notes import parse_script, preservation_changes
from video_gif import gif_info


def sha256(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def audit(pptx: Path, baseline: Path | None = None, script: Path | None = None,
          gif_pages: tuple[int, ...] = ()) -> dict:
    """Check notes equality and required animated media; compare all non-notes parts."""
    with ZipFile(pptx) as archive:
        if archive.testzip() or len(archive.namelist()) != len(set(archive.namelist())):
            raise ValueError("invalid PPT ZIP")
        entries = {name: archive.read(name) for name in archive.namelist()}
    prs = Presentation(pptx)
    report = {"pages": len(prs.slides), "sha256": sha256(pptx), "gif_pages": {},
              "notes_pages": sum(s.has_notes_slide for s in prs.slides)}
    if baseline:
        with ZipFile(baseline) as archive:
            original = {name: archive.read(name) for name in archive.namelist()}
        changes = preservation_changes(original, entries)
        if changes:
            raise ValueError(f"content changed outside notes: {changes}")
        report["content_preserved"] = True
    if script:
        sections, budget = parse_script(script, len(prs.slides))
        for page, slide in enumerate(prs.slides, 1):
            if not slide.has_notes_slide or slide.notes_slide.notes_text_frame.text != sections[page]:
                raise ValueError(f"P{page}: notes differ from Markdown")
        report.update(notes_match_markdown=True, budget_seconds=budget)
    for page in gif_pages:
        if not 1 <= page <= len(prs.slides):
            raise ValueError(f"P{page}: page is outside deck")
        animated = []
        for shape in pictures(prs.slides[page - 1].shapes):
            if shape.image.content_type == "image/gif":
                animated.append(gif_info(shape.image.blob))
        if not animated:
            raise ValueError(f"P{page}: missing animated GIF")
        report["gif_pages"][str(page)] = animated
    return report


def pictures(shapes):
    """Visit pictures actually placed on the page, including nested groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape


def sync_files(root: Path, files: list[Path], host: str, destination: str, timeout: float = 300) -> dict:
    """Copy only requested regular files, never delete; verify remote SHA256 via SSH.

    Uses the host's SSH configuration, agent or terminal-supported authentication.
    Passwords are never accepted as parameters or persisted. Remote filenames are
    shell-quoted.
    """
    if not re.fullmatch(r"[A-Za-z0-9_.@-]+", host) or host.startswith("-"):
        raise ValueError("host must be an SSH alias or user@host")
    if not destination.startswith("/") or any(char in destination for char in "\n\r\\"):
        raise ValueError("remote destination must be an absolute POSIX path")
    root = root.resolve()
    relative = []
    for path in files:
        if path.is_symlink() or not path.is_file():
            raise ValueError(f"not a regular deliverable file: {path}")
        name = path.resolve().relative_to(root).as_posix()
        if any(char in name for char in "\n\r\\"):
            raise ValueError("deliverable filenames cannot contain newlines or backslashes")
        relative.append(name)
    if not relative or len(set(relative)) != len(relative):
        raise ValueError("select at least one file, without duplicates")
    ssh = ["ssh", "-o", "ConnectTimeout=10", host]
    with tempfile.TemporaryDirectory(prefix="ppt-sync-") as directory:
        listing = Path(directory) / "files"
        listing.write_bytes(b"\0".join(name.encode() for name in relative) + b"\0")
        subprocess.run([*ssh, "mkdir -p -- " + shlex.quote(destination)], check=True, timeout=timeout)
        subprocess.run(["rsync", "-a", "--checksum", "--protect-args", "--from0",
                        f"--files-from={listing}", "-e", "ssh -o ConnectTimeout=10",
                        str(root) + "/", f"{host}:{destination.rstrip('/')}/"], check=True, timeout=timeout)
    paths = [str(PurePosixPath(destination) / name) for name in relative]
    command = "sha256sum -- " + " ".join(shlex.quote(path) for path in paths)
    remote = subprocess.run([*ssh, command], capture_output=True, text=True, check=True, timeout=timeout)
    lines = remote.stdout.splitlines()
    expected = {str(PurePosixPath(destination) / name): sha256(root / name) for name in relative}
    actual = {}
    for line in lines:
        digest, name = line.split("  ", 1)
        actual[name] = digest
    if expected != actual:
        raise ValueError("remote SHA256 mismatch; synchronization is not verified")
    return {"host": host, "destination": destination, "files": expected, "sha256_verified": True}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    checker = commands.add_parser("audit")
    checker.add_argument("pptx", type=Path)
    checker.add_argument("--baseline", type=Path)
    checker.add_argument("--script", type=Path)
    checker.add_argument("--gif-pages", type=int, nargs="*", default=[])
    checker.add_argument("--report", type=Path)
    sync = commands.add_parser("sync")
    sync.add_argument("root", type=Path)
    sync.add_argument("--files", nargs="+", required=True, help="paths relative to root")
    sync.add_argument("--host", required=True)
    sync.add_argument("--destination", required=True)
    sync.add_argument("--timeout", type=float, default=300)
    args = parser.parse_args()
    if args.command == "audit":
        report = audit(args.pptx, args.baseline, args.script, tuple(args.gif_pages))
        if args.report:
            args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    else:
        report = sync_files(args.root, [args.root / name for name in args.files], args.host, args.destination, args.timeout)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
