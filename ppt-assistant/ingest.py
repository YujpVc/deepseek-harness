#!/usr/bin/env python3
"""
文件摄取工具 —— 让 PPT 助手「接受」现有 PPT / PDF / 图片文件：
读取 .pptx / .pdf 的内容与结构，输出可读文本供助手分析、重排、转写。
扫描版 PDF（无文字层）自动走 OCR（pdftoppm + tesseract）。

用法：
    python3 ingest.py <file.pptx|file.pdf|file.png|file.jpg>
"""
import sys, os, subprocess, glob, tempfile

def ingest_pptx(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    out = []
    out.append(f"# PPT 摄取: {os.path.basename(path)}")
    out.append(f"页数: {len(prs.slides)}  尺寸: {prs.slide_width/914400:.2f}x{prs.slide_height/914400:.2f} in")
    for i, s in enumerate(prs.slides, 1):
        layout = s.slide_layout.name if s.slide_layout else '?'
        out.append(f"\n--- 第{i}页 (layout: {layout}) ---")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(f"  [文本] {sh.text_frame.text.strip()}")
            elif sh.has_table:
                tbl = sh.table
                out.append(f"  [表格] {len(tbl.rows)}行 x {len(tbl.columns)}列")
                for r in tbl.rows:
                    out.append("     | " + " | ".join(c.text.strip() for c in r.cells))
            elif hasattr(sh, 'has_chart') and sh.has_chart:
                out.append(f"  [图表] {sh.name}")
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                out.append(f"  [图片] {sh.name}")
        if s.has_notes_slide:
            note = s.notes_slide.notes_text_frame.text.strip()
            if note:
                out.append(f"  [备注] {note}")
    return "\n".join(out)

def _tesseract_ocr(image_path):
    """对单张图片 OCR（中英混合）。"""
    r = subprocess.run(
        ["tesseract", image_path, "stdout", "-l", "chi_sim+eng", "--psm", "3"],
        capture_output=True, text=True)
    return (r.stdout or "").strip()

def ingest_image(path):
    txt = _tesseract_ocr(path)
    return f"# 图片 OCR: {os.path.basename(path)}\n{txt}"

def _ocr_scanned_pdf(path):
    """扫描版 PDF：逐页渲染为 PNG 再 OCR。"""
    out = [f"# PDF 摄取(扫描版, 经 OCR): {os.path.basename(path)}"]
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(
            ["pdftoppm", "-png", "-r", "200", path, os.path.join(td, "page")],
            check=True, capture_output=True)
        imgs = sorted(glob.glob(os.path.join(td, "page-*.png")))
        out.append(f"总页数: {len(imgs)}")
        for i, img in enumerate(imgs, 1):
            txt = _tesseract_ocr(img)
            out.append(f"\n--- 第{i}页(OCR) ---\n{txt}")
    return "\n".join(out)

def ingest_pdf(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    pages = []
    total_text = ""
    for i, page in enumerate(r.pages, 1):
        t = page.extract_text() or ""
        total_text += t
        pages.append(f"\n--- 第{i}页 ---\n{t}")
    if total_text.strip():
        return f"# PDF 摄取: {os.path.basename(path)}\n总页数: {len(pages)}" + "".join(pages)
    # 无文字层 → 扫描版，走 OCR
    return _ocr_scanned_pdf(path)

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("用法: python3 ingest.py <file.pptx|file.pdf|file.png|file.jpg>")
        sys.exit(1)
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    if ext == '.pptx':
        print(ingest_pptx(path))
    elif ext == '.pdf':
        print(ingest_pdf(path))
    elif ext in ('.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp'):
        print(ingest_image(path))
    else:
        print(f"不支持的类型: {ext}（仅支持 .pptx / .pdf / 图片）")
        sys.exit(2)
