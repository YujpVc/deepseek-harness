#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
《终末期患者压力性损伤护理专家共识（2026版）》——面向内部医生的共识解读与临床决策简报
专用生成器。16:9，共 20 页。
配色：主色深青绿 #0F766E，辅色暖琥珀 #D97706，正文深灰，白/极浅灰背景。
字体：中文微软雅黑，英文 Arial。页脚统一。标题左侧短色条。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------- 设计令牌 ----------------
TEAL       = RGBColor(0x0F, 0x76, 0x6E)
TEAL_DK    = RGBColor(0x0B, 0x57, 0x52)
TEAL_TINT  = RGBColor(0xE7, 0xF2, 0xF0)
AMBER      = RGBColor(0xD9, 0x77, 0x06)
AMBER_DK   = RGBColor(0xB4, 0x53, 0x09)
AMBER_TINT = RGBColor(0xFB, 0xF1, 0xE2)
INK        = RGBColor(0x1F, 0x29, 0x37)
BODY       = RGBColor(0x3F, 0x46, 0x4F)
MUTED      = RGBColor(0x6B, 0x72, 0x80)
FAINT      = RGBColor(0x9C, 0xA3, 0xAF)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
PANEL      = RGBColor(0xF6, 0xF7, 0xF8)
LINE       = RGBColor(0xE2, 0xE5, 0xE9)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

CN = "微软雅黑"
EN = "Arial"

FOOTER_LEFT = "终末期患者压力性损伤护理专家共识(2026版)｜内部教学"

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


# ---------------- 基础辅助 ----------------
def _set_font(run, size, bold=False, color=BODY, italic=False, latin=EN, ea=CN):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = latin
    rPr = run._r.get_or_add_rPr()
    el = rPr.find(qn('a:ea'))
    if el is None:
        el = rPr.makeelement(qn('a:ea'), {})
        rPr.append(el)
    el.set('typeface', ea)


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.0, round_=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if round_:
        try:
            shp.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.0, wrap=True):
    """paras: list of list-of-runs; run = (text, size, bold, color)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        for (t, sz, bd, col) in runs:
            r = p.add_run()
            r.text = t
            _set_font(r, sz, bd, col)
    return tb


def shape_text(shape, t, size, bold=False, color=WHITE, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = t
    _set_font(r, size, bold, color)


def header(slide, title, module, page):
    box(slide, 0, 0, SLIDE_W, Inches(0.10), fill=TEAL)
    # 章节标记
    if module:
        text(slide, Inches(0.5), Inches(0.30), Inches(11.0), Inches(0.30),
             [[(module, 11, False, TEAL)]])
    # 标题左侧短色条 + 标题
    box(slide, Inches(0.5), Inches(0.62), Inches(0.13), Inches(0.52), fill=AMBER)
    text(slide, Inches(0.78), Inches(0.58), Inches(12.0), Inches(0.62),
         [[(title, 30, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(slide, page)


def footer(slide, page):
    box(slide, Inches(0.5), Inches(7.06), Inches(12.33), Pt(0.8), fill=LINE)
    text(slide, Inches(0.5), Inches(7.14), Inches(10.5), Inches(0.28),
         [[(FOOTER_LEFT, 8.5, False, MUTED)]])
    text(slide, Inches(11.6), Inches(7.14), Inches(1.23), Inches(0.28),
         [[(str(page), 8.5, False, MUTED)]], align=PP_ALIGN.RIGHT)


def doctor_box(slide, x, y, w, h, lines, title="医生视角"):
    box(slide, x, y, w, h, fill=AMBER_TINT, round_=True)
    box(slide, x, y + Inches(0.10), Inches(0.09), h - Inches(0.20), fill=AMBER)
    paras = [[(title, 13, True, AMBER_DK)]]
    for ln in lines:
        paras.append([(ln, 12.5, False, BODY)])
    text(slide, x + Inches(0.26), y + Inches(0.12), w - Inches(0.44), h - Inches(0.24),
         paras, line_spacing=1.05)


def bullet_text(slide, x, y, w, h, items, size=16, gap=0.30, color=BODY, lead_color=TEAL):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            head, body = it
            paras.append([("▪ ", size, True, lead_color), (head, size, True, color),
                          (" " + body, size, False, BODY)])
        else:
            paras.append([("▪ ", size, True, lead_color), (it, size, False, color)])
    text(slide, x, y, w, h, paras, line_spacing=1.12)


def big_stat(slide, x, y, w, num, num_size, label, num_color=TEAL):
    text(slide, x, y, w, Inches(0.6), [[(num, num_size, True, num_color)]])
    text(slide, x, y + Inches(0.62), w, Inches(0.6),
         [[(label, 12, False, BODY)]], line_spacing=1.05)


def set_notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ============================================================
# P1 封面（纯文字）
# ============================================================
def page_cover():
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, SLIDE_W, Inches(0.10), fill=TEAL)
    box(s, Inches(0.7), Inches(1.75), Inches(0.14), Inches(3.1), fill=AMBER)
    text(s, Inches(1.1), Inches(1.85), Inches(11.4), Inches(1.5),
         [[("终末期患者压力性损伤护理专家共识（2026版）", 34, True, INK)]], line_spacing=1.15)
    text(s, Inches(1.1), Inches(3.05), Inches(11.4), Inches(0.6),
         [[("面向内部医生的共识解读与临床决策要点", 20, False, BODY)]])
    box(s, Inches(1.1), Inches(3.95), Inches(2.3), Pt(2.2), fill=AMBER)
    text(s, Inches(1.1), Inches(4.25), Inches(11.4), Inches(0.4),
         [[("中华老年医学杂志 2026年6月 第45卷第6期", 14, False, MUTED)]])
    # 内部教学使用 徽标
    badge = box(s, Inches(1.1), Inches(4.85), Inches(1.9), Inches(0.46), fill=TEAL_TINT,
                line=TEAL, line_w=1.0, round_=True)
    shape_text(badge, "内部教学使用", 13, True, TEAL_DK)
    text(s, Inches(1.1), Inches(6.05), Inches(11.4), Inches(0.4),
         [[("汇报人：____________      日期：____________", 15, False, BODY)]])
    footer(s, 1)
    set_notes(s, "【讲什么】开场点明定位：这是一份面向本院医生的共识解读与临床决策简报，不是护理操作手册。\n"
                 "【强调】说明内部教学用途、来源期刊与期次（中华老年医学杂志 2026年6月 第45卷第6期）。\n"
                 "【追问】本共识与既往压疮/压力性损伤指南是什么关系？适用范围多大？")
    return s


# ============================================================
# P2 目录
# ============================================================
def page_toc():
    s = prs.slides.add_slide(BLANK)
    header(s, "目录", "", 2)
    modules = [
        ("01", "背景与定义", "终末期与压力性损伤（PI）的界定"),
        ("02", "共识如何制订", "注册 · 流程 · 检索 · 证据分级"),
        ("03", "核心理念与总览", "从「愈合」转向「舒适与尊严」"),
        ("04", "医生临床决策要点", "评估 · 预防 · 治疗 · 转介 · 质控"),
        ("05", "讨论与展望", "证据局限与未来研究方向"),
        ("06", "来源声明", "出处 · 基金 · 利益冲突"),
    ]
    x0, y0 = Inches(0.5), Inches(1.55)
    cw, ch = Inches(6.0), Inches(1.52)
    gx, gy = Inches(0.33), Inches(0.28)
    for i, (num, name, desc) in enumerate(modules):
        col, row = i % 2, i // 2
        x = x0 + col * (cw + gx)
        y = y0 + row * (ch + gy)
        box(s, x, y, cw, ch, fill=PANEL, round_=True)
        box(s, x, y, Inches(0.09), ch, fill=TEAL)
        text(s, x + Inches(0.32), y + Inches(0.22), Inches(1.0), Inches(0.6),
             [[(num, 24, True, TEAL)]])
        text(s, x + Inches(1.35), y + Inches(0.20), Inches(4.4), Inches(0.5),
             [[(name, 19, True, INK)]])
        text(s, x + Inches(1.35), y + Inches(0.78), Inches(4.4), Inches(0.5),
             [[(desc, 13, False, MUTED)]])
    set_notes(s, "【讲什么】用约30秒介绍六大模块与整体叙事主线：为什么重要 → 怎么科学制订 → 医生该做什么 → 证据边界。\n"
                 "【强调】医生最关心的内容集中在模块四「临床决策要点」。\n"
                 "【追问】（无）")
    return s


# ============================================================
# P3 背景与问题（左数据带 + 右为什么医生要关注）
# ============================================================
def page_background():
    s = prs.slides.add_slide(BLANK)
    header(s, "背景与问题", "模块一 · 背景与定义", 3)
    # 左侧数据带
    lx, lw = Inches(0.5), Inches(4.5)
    box(s, lx, Inches(1.5), lw, Inches(5.35), fill=TEAL, round_=True)
    text(s, lx + Inches(0.35), Inches(1.72), lw - Inches(0.7), Inches(0.4),
         [[("终末期患者压力性损伤（PI）数据带", 14, True, WHITE)]])
    stats = [
        ("12.4%~34.1%", "压力性损伤患病率"),
        ("11.7%~26.5%", "压力性损伤发生率"),
        ("≤6 个月", "终末期预期生存期"),
        ("翻身困难 / 营养不良 / 持续镇静镇痛", "显著增加 PI 风险的高危因素"),
    ]
    sy = Inches(2.35)
    for num, lab in stats:
        text(s, lx + Inches(0.35), sy, lw - Inches(0.7), Inches(0.6),
             [[(num, 25, True, WHITE)]])
        text(s, lx + Inches(0.35), sy + Inches(0.62), lw - Inches(0.7), Inches(0.55),
             [[(lab, 11.5, False, RGBColor(0xDD, 0xEC, 0xEA))]], line_spacing=1.0)
        sy += Inches(1.14)
    # 右侧：为什么医生要关注
    rx = Inches(5.3)
    rw = Inches(7.53)
    text(s, rx, Inches(1.52), rw, Inches(0.45),
         [[("为什么医生要关注", 18, True, TEAL_DK)]])
    box(s, rx, Inches(1.98), Inches(0.5), Pt(2.4), fill=AMBER)
    bullet_text(s, rx, Inches(2.22), rw, Inches(3.2), [
        ("定义：", "终末期指疾病进程不可逆、预期生存≤6个月；人群正由癌症为主转向老年及慢病为主。"),
        ("后果：", "PI 可加重疼痛、睡眠障碍、焦虑、抑郁，降低生存质量。"),
        ("缺口：", "现有指南及共识对终末期患者关注不足。"),
    ], size=15, gap=0.34)
    doctor_box(s, rx, Inches(4.85), rw, Inches(1.95), [
        "明确治疗边界，避免过度干预；",
        "组织多学科团队（MDT）并承担转介决策；",
        "在「积极预防」与「避免过度干预」之间权衡。",
    ])
    set_notes(s, "【讲什么】终末期定义（预期生存≤6个月）、人群转向老年/慢病、PI患病率12.4%~34.1%、发生率11.7%~26.5%。\n"
                 "【强调】PI不仅影响皮肤，还加重疼痛/睡眠障碍/焦虑抑郁、降低生存质量；现有指南对终末期关注不足，是制订本共识的原因。\n"
                 "【追问】患病率区间为何这么宽（12.4%~34.1%）？终末期患者 PI 是否不可避免？")
    return s


# ============================================================
# P4 共识制订（一）组织与流程
# ============================================================
def page_method1():
    s = prs.slides.add_slide(BLANK)
    header(s, "共识制订：组织与流程", "模块二 · 共识如何制订", 4)
    # 注册 + 依据
    text(s, Inches(0.5), Inches(1.42), Inches(12.33), Inches(0.35),
         [[("注册：", 13, True, TEAL_DK),
           ("国际实践指南注册与透明化平台（双语）｜PREPARE-2026CN245", 13, False, BODY)]])
    text(s, Inches(0.5), Inches(1.86), Inches(12.33), Inches(0.35),
         [[("依据：", 13, True, TEAL_DK),
           ("《中国制订/修订临床诊疗指南的指导原则（2022版）》｜参考 RIGHT 报告条目", 13, False, BODY)]])
    # 时间线
    text(s, Inches(0.5), Inches(2.42), Inches(6.0), Inches(0.35),
         [[("关键时间线", 16, True, INK)]])
    tl_y = Inches(3.15)
    box(s, Inches(0.9), tl_y + Inches(0.24), Inches(11.6), Pt(2.2), fill=TEAL_TINT)
    nodes = [("2025年11月", "项目启动"), ("2026年1月", "补充检索"), ("2026年4月", "成稿")]
    for i, (d, ev) in enumerate(nodes):
        cx = Inches(1.3) + i * Inches(5.3)
        box(s, cx, tl_y + Inches(0.06), Inches(0.36), Inches(0.36), fill=TEAL,
            round_=True)
        text(s, cx - Inches(0.7), tl_y - Inches(0.45), Inches(2.0), Inches(0.4),
             [[(d, 14, True, TEAL_DK)]], align=PP_ALIGN.CENTER)
        text(s, cx - Inches(0.7), tl_y + Inches(0.55), Inches(2.0), Inches(0.4),
             [[(ev, 13, False, BODY)]], align=PP_ALIGN.CENTER)
    # 委员会构成数字条
    text(s, Inches(0.5), Inches(4.55), Inches(6.0), Inches(0.35),
         [[("委员会构成（共 52 人）", 16, True, INK)]])
    groups = [("指导专家组", "9"), ("函询专家组", "16"), ("外部评审专家组", "13"),
              ("共识编写组", "10"), ("秘书组", "4")]
    card_w = Inches(2.32)
    gap = Inches(0.185)
    x0 = Inches(0.5)
    for i, (nm, num) in enumerate(groups):
        x = x0 + i * (card_w + gap)
        box(s, x, Inches(5.0), card_w, Inches(1.78), fill=TEAL_TINT, round_=True)
        text(s, x + Inches(0.2), Inches(5.18), card_w - Inches(0.4), Inches(0.6),
             [[(num, 30, True, TEAL_DK)]])
        text(s, x + Inches(0.2), Inches(5.92), card_w - Inches(0.4), Inches(0.7),
             [[(nm, 12.5, False, BODY)]], line_spacing=1.0)
    set_notes(s, "【讲什么】共识已双语注册（PREPARE-2026CN245），依据2022版指导原则、参考RIGHT；五类委员会共52人；时间线2025.11启动→2026.1补充检索→2026.4成稿。\n"
                 "【强调】方法学规范（注册、指导原则、多类委员会、外部评审）是共识可信度的基础。\n"
                 "【追问】函询专家组与外部评审专家组分工有何不同？PREPARE注册号意味着什么？")
    return s


# ============================================================
# P5 共识制订（二）问题与检索
# ============================================================
def page_method2():
    s = prs.slides.add_slide(BLANK)
    header(s, "共识制订：问题与检索", "模块二 · 共识如何制订", 5)
    # 左：漏斗
    lx = Inches(0.5)
    text(s, lx, Inches(1.42), Inches(5.6), Inches(0.35),
         [[("从问题到主题（PIPOST 框架）", 15, True, TEAL_DK)]])
    fun = [("初始问题  27 个", 5.4, TEAL), ("两轮线下投票", 5.4, AMBER),
           ("临床问题  20 个", 4.0, TEAL), ("归纳", 4.0, AMBER),
           ("9 大主题", 2.8, TEAL_DK)]
    fy = Inches(1.88)
    fcx = lx + Inches(2.7)
    for label, w, col in fun:
        x = fcx - Inches(w / 2)
        box(s, x, fy, Inches(w), Inches(0.52), fill=col, round_=True)
        shape_text(s.shapes[-1], label, 14, True, WHITE)
        fy += Inches(0.74)
    text(s, lx, Inches(5.6), Inches(5.6), Inches(0.75),
         [[("前期文献梳理 + 对 ICU 及安宁疗护病房患者、照顾者、临床护士进行深度访谈，"
            "采用 PIPOST 框架构建问题。", 11.5, False, MUTED)]], line_spacing=1.1)
    # 右：检索与纳入
    rx = Inches(6.45)
    rw = Inches(6.4)
    text(s, rx, Inches(1.42), rw, Inches(0.35),
         [[("检索来源与纳入", 15, True, TEAL_DK)]])
    text(s, rx, Inches(1.86), rw, Inches(1.6),
         [[("数据库：", 12, True, BODY),
           ("PubMed · Embase · Web of Science · Cochrane Library · Scopus · CINAHL · SinoMed · 中国知网 · 万方 · 维普", 12, False, BODY)]],
         line_spacing=1.15)
    text(s, rx, Inches(2.85), rw, Inches(1.6),
         [[("专业组织：", 12, True, BODY),
           ("NICE · GIN · SIGN · RNAO · UpToDate · BMJ Best Practice · JBI · Wounds International · EWMA · WUWHS · WOCN · EPUAP · NPIAP", 12, False, BODY)]],
         line_spacing=1.15)
    text(s, rx, Inches(3.95), rw, Inches(0.4),
         [[("初检 527 篇  →  最终纳入 30 篇", 14, True, TEAL_DK)]])
    # 30篇构成横向堆叠条
    bar_y = Inches(4.5)
    text(s, rx, bar_y - Inches(0.32), rw, Inches(0.3),
         [[("30 篇文献构成", 12.5, True, BODY)]])
    segs = [("指南 4", 4, TEAL_DK), ("专家共识 7", 7, TEAL),
            ("系统评价 3", 3, AMBER), ("原始研究 16", 16, AMBER_DK)]
    total_w = 6.0
    sx = rx
    seg_x = sx
    for label, n, col in segs:
        w = Inches(total_w * n / 30.0)
        box(s, seg_x, bar_y, w, Inches(0.5), fill=col)
        if n >= 4:
            shape_text(s.shapes[-1], label, 12, True, WHITE)
        seg_x += w
    text(s, rx, bar_y + Inches(0.66), rw, Inches(0.5),
         [[("指南 4 · 专家共识 7 · 系统评价 3 · 原始研究 16", 11.5, False, MUTED)]])
    # 9大主题一览
    themes = ["多学科团队", "照护目标", "患者评估", "预防", "治疗",
              "症状管理", "患者教育", "转介时机", "质量控制"]
    text(s, rx, Inches(5.5), rw, Inches(0.3), [[("9 大主题", 12.5, True, BODY)]])
    th = text(s, rx, Inches(5.82), rw, Inches(1.0),
              [[(" · ".join(themes), 11, False, BODY)]], line_spacing=1.15)
    set_notes(s, "【讲什么】PIPOST框架构建问题；初始27题→两轮线下投票→20题，覆盖9大主题；多数据库+专业组织网站检索；初检527篇→最终30篇。\n"
                 "【强调】证据来源广度与专业组织指南检索保证了推荐的可信度；30篇构成（指南4/共识7/系统评价3/原始研究16）。\n"
                 "【追问】为何最终只纳入30篇？9大主题如何覆盖医生最关心的临床问题？")
    return s


# ============================================================
# P6 共识制订（三）证据与推荐
# ============================================================
def page_method3():
    s = prs.slides.add_slide(BLANK)
    header(s, "共识制订：证据与推荐", "模块二 · 共识如何制订", 6)
    # 四步流程
    steps = [
        ("01", "质量评价", "指南用 AGREEⅡ；专家共识用 JBI 标准；系统评价/Meta 分析用 AMSTAR2；原始研究用 JBI 工具"),
        ("02", "双人独立", "文献筛选、质量评价、证据提取均由两人独立完成，分歧时咨询第三人"),
        ("03", "证据分级", "按 JBI 2014 版证据预分级系统分为 1~5 级"),
        ("04", "形成推荐", "结合可行性、适宜性、意义和有效性形成推荐强度（强弱分级）"),
    ]
    cw = Inches(2.95)
    gap = Inches(0.18)
    x0 = Inches(0.5)
    for i, (n, h, t) in enumerate(steps):
        x = x0 + i * (cw + gap)
        box(s, x, Inches(1.45), cw, Inches(2.45), fill=PANEL, round_=True)
        box(s, x, Inches(1.45), cw, Inches(0.16), fill=TEAL if i % 2 == 0 else AMBER)
        text(s, x + Inches(0.2), Inches(1.72), cw - Inches(0.4), Inches(0.5),
             [[(n, 22, True, AMBER if i % 2 == 0 else TEAL)]])
        text(s, x + Inches(0.2), Inches(2.2), cw - Inches(0.4), Inches(0.4),
             [[(h, 15, True, INK)]])
        text(s, x + Inches(0.2), Inches(2.66), cw - Inches(0.4), Inches(1.2),
             [[(t, 11.5, False, BODY)]], line_spacing=1.12)
    # 关键数据栏（德尔菲 + 外部会审）
    text(s, Inches(0.5), Inches(4.05), Inches(6.0), Inches(0.35),
         [[("德尔菲函询与外部会审（关键数据）", 16, True, INK)]])
    kpis = [
        ("16 名", "国内资深专家"),
        ("100%", "函询回收率"),
        ("0.96", "专家权威系数"),
        ("30 个", "保留条目"),
        ("4.63~5 分", "重要性均值"),
        ("93.75%~100%", "认同度"),
        ("0~0.17", "变异系数"),
        ("≥75%", "外部会审确定强弱阈值"),
    ]
    cw = Inches(2.95)
    chh = Inches(0.72)
    gx = Inches(0.17)
    gy = Inches(0.14)
    kx0 = Inches(0.5)
    ky0 = Inches(4.55)
    for i, (num, lab) in enumerate(kpis):
        col, row = i % 4, i // 4
        x = kx0 + col * (cw + gx)
        y = ky0 + row * (chh + gy)
        box(s, x, y, cw, chh, fill=TEAL_TINT, round_=True)
        text(s, x + Inches(0.16), y + Inches(0.05), cw - Inches(0.3), Inches(0.36),
             [[(num, 16, True, TEAL_DK)]])
        text(s, x + Inches(0.16), y + Inches(0.42), cw - Inches(0.3), Inches(0.28),
             [[(lab, 10, False, BODY)]])
    # 结果
    box(s, Inches(0.5), Inches(6.28), Inches(12.33), Inches(0.62), fill=AMBER_TINT, round_=True)
    text(s, Inches(0.75), Inches(6.38), Inches(11.9), Inches(0.45),
         [[("最终形成 29 条推荐意见", 15, True, AMBER_DK),
           ("　（13 名外部专家会审，每项推荐 ≥75% 专家同意后确定强弱）", 12.5, False, BODY)]])
    set_notes(s, "【讲什么】质量评价工具（AGREEⅡ/AMSTAR2/JBI）、双人独立+第三人仲裁、JBI 1~5级、形成推荐强度；两轮德尔菲（权威系数0.96、回收率100%）+外部会审≥75%确定强弱，最终29条。\n"
                 "【强调】过程严谨、专家权威度高（0.96），变异系数0~0.17说明专家意见较集中。\n"
                 "【追问】德尔菲「变异系数0~0.17」说明什么？强弱推荐在临床如何落地执行？")
    return s


# ============================================================
# P7 定义、适用范围与使用边界
# ============================================================
def page_definition():
    s = prs.slides.add_slide(BLANK)
    header(s, "定义与适用范围", "模块一 · 背景与定义", 7)
    # 上半区：定义
    text(s, Inches(0.5), Inches(1.42), Inches(6.0), Inches(0.35),
         [[("核心定义", 16, True, TEAL_DK)]])
    box(s, Inches(0.5), Inches(1.82), Inches(12.33), Inches(1.85), fill=PANEL, round_=True)
    text(s, Inches(0.8), Inches(2.0), Inches(11.8), Inches(1.6), [
        [("终末期患者：", 14, True, TEAL_DK),
         ("因严重疾病、伤害或其他原因无法恢复，且有明确证据表明疾病进程无法逆转、预期生存期在 6 个月以内的患者。", 14, False, BODY)],
        [("压力性损伤（PI, pressure injury）：", 14, True, TEAL_DK),
         ("皮肤和/或皮下组织的局限性损伤，由压力或压力合并剪切力所致，常发生于骨突部位或与医疗器械及其他物品相接触处。", 14, False, BODY)],
    ], line_spacing=1.18)
    # 下半区：适用与边界
    text(s, Inches(0.5), Inches(3.85), Inches(6.0), Inches(0.35),
         [[("适用范围与使用边界", 16, True, TEAL_DK)]])
    col = [
        ("适用对象", ["各级医院从事终末期患者照护的临床医师、护士、医技人员、护理员及科研教学人员"]),
        ("目标人群", ["存在压力性损伤风险或已患病的终末期患者"]),
        ("使用边界", ["复杂患者应在多学科团队支持下制定方案", "本共识为指导工具，不能替代专业人员的临床决策"]),
    ]
    cw = Inches(3.95)
    gap = Inches(0.24)
    x0 = Inches(0.5)
    for i, (h, pts) in enumerate(col):
        x = x0 + i * (cw + gap)
        box(s, x, Inches(4.25), cw, Inches(2.55), fill=WHITE, line=LINE, line_w=1.0, round_=True)
        box(s, x, Inches(4.25), cw, Inches(0.14), fill=TEAL)
        text(s, x + Inches(0.22), Inches(4.52), cw - Inches(0.44), Inches(0.4),
             [[(h, 15, True, INK)]])
        paras = [[("▪ ", 13, True, TEAL), (p, 13, False, BODY)] for p in pts]
        text(s, x + Inches(0.22), Inches(4.98), cw - Inches(0.44), Inches(1.7),
             paras, line_spacing=1.12)
    set_notes(s, "【讲什么】终末期患者与PI的明确定义、适用对象/目标人群、使用边界。\n"
                 "【强调】边界是关键——共识是「指导工具」，不能替代专业人员的临床决策；复杂患者需MDT支持。\n"
                 "【追问】「预期生存6个月」在临床如何判断？共识对非终末期患者是否适用？")
    return s


# ============================================================
# P8 核心理念与29条总览
# ============================================================
def page_overview():
    s = prs.slides.add_slide(BLANK)
    header(s, "核心理念与总览", "模块三 · 核心理念与总览", 8)
    # 核心理念横幅
    box(s, Inches(0.5), Inches(1.42), Inches(12.33), Inches(1.0), fill=TEAL, round_=True)
    text(s, Inches(0.85), Inches(1.52), Inches(11.6), Inches(0.8), [
        [("核心理念：", 17, True, WHITE),
         ("把目标从「愈合」转向「舒适与尊严」", 17, True, WHITE)],
        [("终末期照护不以「不发生 PI 或促进愈合」为唯一目标", 12.5, False, RGBColor(0xDD, 0xEC, 0xEA))],
    ], line_spacing=1.15)
    # 9大主题矩阵（3x3）
    themes = [
        ("多学科团队", "1 条"), ("照护目标", "1 条"),
        ("患者评估", "5 条 · 风险3+皮肤2"), ("预防", "5 条"),
        ("治疗", "2 条"), ("症状管理", "11 条"),
        ("患者教育", "2 条"), ("转介", "1 条"),
        ("质量控制", "1 条"),
    ]
    cw = Inches(3.95)
    ch = Inches(1.42)
    gx = Inches(0.24)
    gy = Inches(0.16)
    x0 = Inches(0.5)
    y0 = Inches(2.62)
    for i, (nm, cnt) in enumerate(themes):
        col, row = i % 3, i // 3
        x = x0 + col * (cw + gx)
        y = y0 + row * (ch + gy)
        big = (nm == "症状管理")
        box(s, x, y, cw, ch, fill=AMBER_TINT if big else PANEL,
            line=AMBER if big else LINE, line_w=1.2 if big else 0.8, round_=True)
        text(s, x + Inches(0.22), y + Inches(0.18), cw - Inches(0.44), Inches(0.5),
             [[(nm, 15.5, True, AMBER_DK if big else INK)]])
        text(s, x + Inches(0.22), y + Inches(0.78), cw - Inches(0.44), Inches(0.5),
             [[(cnt, 13, True, TEAL if not big else AMBER_DK)]])
    text(s, Inches(0.5), Inches(6.62), Inches(12.33), Inches(0.4),
         [[("29 条推荐意见按 9 大主题分布，其中「症状管理」占比最高（11 条）。", 12.5, False, MUTED)]])
    set_notes(s, "【讲什么】核心理念「愈合→舒适与尊严」；29条推荐按9大主题分布，症状管理11条占比最高。\n"
                 "【强调】理念转向是理解全部推荐的前提——终末期不以「不发生PI/愈合」为唯一目标。\n"
                 "【追问】为什么「症状管理」占11条（近四成）？这与终末期照护目标如何对应？")
    return s


# ============================================================
# P9 多学科团队与照护目标
# ============================================================
def page_mdt():
    s = prs.slides.add_slide(BLANK)
    header(s, "多学科团队与照护目标", "模块三 · 核心理念与总览", 9)
    # 左：团队放射图
    text(s, Inches(0.5), Inches(1.42), Inches(5.6), Inches(0.35),
         [[("多学科团队（MDT）", 16, True, TEAL_DK)]])
    cx = Inches(3.2)
    cy = Inches(3.5)
    box(s, cx - Inches(0.95), cy - Inches(0.5), Inches(1.9), Inches(1.0),
        fill=TEAL, round_=True)
    text(s, cx - Inches(0.95), cy - Inches(0.5), Inches(1.9), Inches(1.0), [
        [("多学科团队", 14, True, WHITE)],
        [("MDT", 11, False, RGBColor(0xDD, 0xEC, 0xEA))],
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    members = ["伤口护理专家", "医生", "护士", "营养师", "物理治疗师",
               "心理治疗师", "药剂师", "社会工作者"]
    import math
    R = Inches(1.55)
    for i, m in enumerate(members):
        ang = math.radians(-90 + i * 45)
        mx = cx + R * math.cos(ang)
        my = cy + R * math.sin(ang)
        w = Inches(1.5)
        h = Inches(0.5)
        b = box(s, mx - w / 2, my - h / 2, w, h, fill=TEAL_TINT, line=TEAL, line_w=0.8, round_=True)
        shape_text(b, m, 11.5, False, TEAL_DK)
    text(s, Inches(0.5), Inches(5.5), Inches(5.6), Inches(1.5), [
        [("资源有限时至少保留：", 12, True, AMBER_DK),
         ("伤口护理专家、专科医护人员、营养师、药剂师及社会工作者。", 12, False, BODY)],
        [("团队中至少有一名", 12, True, BODY),
         ("有资质的伤口护理师。", 12, True, AMBER_DK)],
    ], line_spacing=1.15)
    # 右：目标阶梯
    rx = Inches(6.3)
    rw = Inches(6.5)
    text(s, rx, Inches(1.42), rw, Inches(0.35), [[("照护目标", 16, True, TEAL_DK)]])
    steps = [
        ("个体化目标", "符合患者价值观与意愿，并参考照护者意见", TEAL, WHITE),
        ("聚焦核心", "舒适体验 · 症状控制 · 生活质量", TEAL_DK, WHITE),
        ("非唯一目标", "不以「不发生 PI 或促进愈合」为唯一追求", AMBER, WHITE),
    ]
    sw = Inches(5.9)
    sy = Inches(1.95)
    for i, (h, t, col, tc) in enumerate(steps):
        w = sw - i * Inches(0.5)
        x = rx + (sw - w) / 2
        box(s, x, sy, w, Inches(0.78), fill=col, round_=True)
        text(s, x + Inches(0.25), sy, w - Inches(0.5), Inches(0.78), [
            [(h + "　", 14, True, tc), (t, 11.5, False, tc)],
        ], anchor=MSO_ANCHOR.MIDDLE)
        sy += Inches(0.92)
    box(s, rx, Inches(4.85), rw, Inches(0.62), fill=AMBER_TINT, round_=True)
    text(s, rx + Inches(0.25), Inches(4.94), rw - Inches(0.5), Inches(0.45),
         [[("若符合患者意愿，仍可把「治愈」作为目标。", 13, True, AMBER_DK)]])
    doctor_box(s, rx, Inches(5.62), rw, Inches(1.2), [
        "明确自己在原发病管理、药物决策、转介中的角色；",
        "参与家庭会议，与护理团队对齐照护目标。",
    ])
    set_notes(s, "【讲什么】MDT八类角色构成、资源有限时的最小团队、至少1名有资质伤口护理师；照护目标个体化，聚焦舒适/症状/生活质量。\n"
                 "【强调】医生角色明确——原发病管理、药物决策、转介；参与家庭会议对齐目标。\n"
                 "【追问】本院MDT目前缺哪些角色？「治愈」作为目标在什么情形仍适用？")
    return s


# ============================================================
# P10 风险评估与皮肤评估
# ============================================================
def page_assessment():
    s = prs.slides.add_slide(BLANK)
    header(s, "风险评估与皮肤评估", "模块四 · 医生临床决策要点", 10)
    # 风险因素三分类
    text(s, Inches(0.5), Inches(1.4), Inches(6.0), Inches(0.35),
         [[("风险因素三分类", 16, True, TEAL_DK)]])
    cats = [
        ("全身因素", "年龄、营养不良、慢性合并症、PI 病史、严重低蛋白、电解质失衡、灌注/氧合不足、活动能力下降、体温异常、重度水肿、器官功能衰竭、疼痛及精神心理因素"),
        ("局部因素", "持续受压、剪切力、潮湿等"),
        ("治疗因素", "抗癌治疗、住院时间、血管活性药/镇静剂/免疫抑制剂等特殊药物、医疗器械使用、治疗依从性等"),
    ]
    cw = Inches(3.95)
    gap = Inches(0.24)
    x0 = Inches(0.5)
    for i, (h, t) in enumerate(cats):
        x = x0 + i * (cw + gap)
        box(s, x, Inches(1.76), cw, Inches(2.15), fill=PANEL, round_=True)
        box(s, x, Inches(1.76), cw, Inches(0.14), fill=TEAL)
        text(s, x + Inches(0.2), Inches(2.0), cw - Inches(0.4), Inches(0.4),
             [[(h, 14, True, INK)]])
        text(s, x + Inches(0.2), Inches(2.44), cw - Inches(0.4), Inches(1.4),
             [[(t, 10.5, False, BODY)]], line_spacing=1.12)
    # 左下：评估工具表
    text(s, Inches(0.5), Inches(4.1), Inches(6.0), Inches(0.35),
         [[("评估工具选择", 15, True, INK)]])
    rows = [["住院患者", "Waterlow 量表"],
            ["居家 / 非住院", "Braden 量表"],
            ["其他可选", "PPS · Hunter's Hill 量表 · HoRT 量表"]]
    tbl_shape = s.shapes.add_table(4, 2, Inches(0.5), Inches(4.52), Inches(5.7), Inches(1.7))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(2.2)
    tbl.columns[1].width = Inches(3.5)
    hdrs = ["场景", "推荐工具"]
    data = [hdrs] + rows
    for ri, r in enumerate(data):
        for ci, v in enumerate(r):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = TEAL
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if ri % 2 else PANEL
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = v
            _set_font(run, 12, ri == 0, WHITE if ri == 0 else BODY)
    # 右下：评估频率
    rx = Inches(6.6)
    rw = Inches(6.23)
    text(s, rx, Inches(4.1), rw, Inches(0.35), [[("评估频率与重点", 15, True, INK)]])
    freq = [
        ("入院 · 转科 · 病情变化", "结构化评估"),
        ("有风险者", "至少每 4~7 天复评；病情波动随时评估"),
        ("皮肤评估", "骨突处 · 失禁区 · 头皮毛发下 · 器械接触部位；住院至少每日两次"),
    ]
    fy = Inches(4.5)
    for h, t in freq:
        box(s, rx, fy, rw, Inches(0.58), fill=WHITE, line=LINE, line_w=0.8, round_=True)
        box(s, rx, fy, Inches(0.08), Inches(0.58), fill=AMBER)
        text(s, rx + Inches(0.22), fy, rw - Inches(0.4), Inches(0.58), [
            [(h + "　", 12.5, True, INK), (t, 12, False, BODY)],
        ], anchor=MSO_ANCHOR.MIDDLE)
        fy += Inches(0.68)
    doctor_box(s, Inches(0.5), Inches(6.45), Inches(12.33), Inches(0.55), [
        "关注高风险药物与治疗决策，复核风险评估记录；发现高风险时启动预防并组织多学科会诊。",
    ])
    set_notes(s, "【讲什么】风险因素三分类（全身/局部/治疗）；工具选择（住院Waterlow、居家Braden、其他PPS等）；评估频率（4~7天复评、皮肤每日至少两次）。\n"
                 "【强调】医生重点关注高风险药物与治疗决策（血管活性药、镇静剂、抗癌治疗等），复核评估记录。\n"
                 "【追问】为何住院与居家工具不同？4~7天复评的循证依据？")
    return s


# ============================================================
# P11 预防一：敷料与营养
# ============================================================
def page_prevention1():
    s = prs.slides.add_slide(BLANK)
    header(s, "预防一：敷料与营养", "模块四 · 医生临床决策要点", 11)
    lx = Inches(0.5)
    rw = Inches(6.0)
    rx = Inches(6.75)
    # 左：敷料
    box(s, lx, Inches(1.45), rw, Inches(2.4), fill=PANEL, round_=True)
    box(s, lx, Inches(1.45), rw, Inches(0.14), fill=TEAL)
    text(s, lx + Inches(0.25), Inches(1.72), rw - Inches(0.5), Inches(0.4),
         [[("预防性敷料", 16, True, INK)]])
    bullet_text(s, lx + Inches(0.25), Inches(2.2), rw - Inches(0.5), Inches(1.55), [
        "高危患者在骶尾部及足跟应用预防性敷料",
        "其他持续受压部位酌情使用",
        "所有预防性敷料应在确认存在风险后使用",
    ], size=13, gap=0.26)
    # 右：营养
    box(s, rx, Inches(1.45), rw, Inches(4.6), fill=PANEL, round_=True)
    box(s, rx, Inches(1.45), rw, Inches(0.14), fill=AMBER)
    text(s, rx + Inches(0.25), Inches(1.72), rw - Inches(0.5), Inches(0.4),
         [[("营养支持", 16, True, INK)]])
    bullet_text(s, rx + Inches(0.25), Inches(2.2), rw - Inches(0.5), Inches(3.7), [
        ("个体化：", "根据患者意愿及生存期限考虑；优先经口"),
        ("经口不足：", "充分评估风险与获益，与患者及家属商议替代方案，基于耐受度与舒适度实施"),
        ("管饲喂养：", "不一定能降低吸入性肺炎风险、提高生活质量、降低死亡率或减轻疼痛"),
        ("晚期痴呆或预期生存≤1个月：", "不推荐采用饮食摄入以外的替代营养途径"),
        ("不宜", "把能量或蛋白达标作为目标"),
    ], size=13, gap=0.30)
    # 医生视角
    doctor_box(s, lx, Inches(4.05), rw, Inches(2.0), [
        "营养方案是医嘱决策点，应与家属充分沟通；",
        "避免过度营养干预；",
        "尊重患者意愿与生存期限，权衡获益与负担。",
    ])
    set_notes(s, "【讲什么】预防性敷料（骶尾/足跟）；营养支持（优先经口、个体化）；管饲不必然获益；晚期痴呆/预期生存≤1月不推荐替代营养途径；不以能量蛋白达标为目标。\n"
                 "【强调】医生在营养医嘱决策中应避免过度营养干预，与家属充分沟通。\n"
                 "【追问】管饲喂养的循证证据强度如何？何时应停用替代营养途径？")
    return s


# ============================================================
# P12 预防二：体位管理
# ============================================================
def page_prevention2():
    s = prs.slides.add_slide(BLANK)
    header(s, "预防二：体位管理", "模块四 · 医生临床决策要点", 12)
    cards = [
        ("个性化翻身计划", "根据照护目标、舒适度、耐受性及患者意愿制定", TEAL),
        ("无法耐受翻身", "以患者耐受为限度进行小幅度体位调整，或仅移动病情允许的身体部位", AMBER),
        ("减压辅具", "结合减压坐垫/床垫适当调整翻身间隔，提升舒适度", TEAL),
        ("持续沟通", "与患者和家属持续沟通，不机械执行翻身频率", TEAL_DK),
    ]
    cw = Inches(2.95)
    gap = Inches(0.18)
    x0 = Inches(0.5)
    for i, (h, t, col) in enumerate(cards):
        x = x0 + i * (cw + gap)
        box(s, x, Inches(1.55), cw, Inches(2.6), fill=PANEL, round_=True)
        box(s, x, Inches(1.55), cw, Inches(0.16), fill=col)
        text(s, x + Inches(0.2), Inches(1.85), cw - Inches(0.4), Inches(0.75),
             [[(h, 15, True, INK)]], line_spacing=1.0)
        text(s, x + Inches(0.2), Inches(2.62), cw - Inches(0.4), Inches(1.45),
             [[(t, 12, False, BODY)]], line_spacing=1.15)
    # 时间线：翻身节奏示意
    text(s, Inches(0.5), Inches(4.45), Inches(6.0), Inches(0.35),
         [[("体位调整节奏（示意）", 15, True, INK)]])
    tl_y = Inches(5.15)
    box(s, Inches(0.9), tl_y + Inches(0.22), Inches(11.6), Pt(2.2), fill=TEAL_TINT)
    nodes = [("评估", "照护目标·耐受度"), ("计划", "个体化翻身计划"),
             ("调整", "减压坐垫/床垫"), ("沟通", "家属持续沟通")]
    for i, (h, t) in enumerate(nodes):
        cx = Inches(1.3) + i * Inches(3.5)
        box(s, cx, tl_y + Inches(0.05), Inches(0.34), Inches(0.34), fill=AMBER, round_=True)
        text(s, cx - Inches(1.0), tl_y - Inches(0.42), Inches(2.35), Inches(0.4),
             [[(h, 14, True, TEAL_DK)]], align=PP_ALIGN.CENTER)
        text(s, cx - Inches(1.0), tl_y + Inches(0.52), Inches(2.35), Inches(0.4),
             [[(t, 11, False, BODY)]], align=PP_ALIGN.CENTER)
    doctor_box(s, Inches(0.5), Inches(6.15), Inches(12.33), Inches(0.8), [
        "支持护理团队灵活调整方案；避免不必要的检查或操作打断患者的舒适护理。",
    ])
    set_notes(s, "【讲什么】个性化翻身计划、不耐受者小幅度调整、结合减压坐垫/床垫、不机械执行翻身频率。\n"
                 "【强调】医生应支持护理团队灵活调整，避免不必要操作打断舒适护理。\n"
                 "【追问】翻身频率到底多久一次？减压床垫如何选择？")
    return s


# ============================================================
# P13 症状管理一：清创与异味
# ============================================================
def page_symptom1():
    s = prs.slides.add_slide(BLANK)
    header(s, "症状管理一：清创与异味", "模块四 · 医生临床决策要点", 13)
    text(s, Inches(0.5), Inches(1.4), Inches(6.0), Inches(0.35),
         [[("清洗与清创原则", 15, True, TEAL_DK)]])
    bullet_text(s, Inches(0.5), Inches(1.8), Inches(12.3), Inches(1.15), [
        ("轻柔清洗：", "伤口清洗尽可能轻柔，避免额外组织损伤。"),
        ("谨慎清创：", "对潜在血管疾病或干性坏疽者避免激进清创。"),
    ], size=14, gap=0.3)
    # 异味处理路径图（横向流程）
    text(s, Inches(0.5), Inches(3.05), Inches(6.0), Inches(0.35),
         [[("异味处理路径", 15, True, TEAL_DK)]])
    steps = [
        ("① 优先消除病因", "按培养结果应用敏感抗生素 · 局部抗菌溶液冲洗 · 抗菌敷料控制感染", TEAL),
        ("② 难控 / 临终", "组织坏死及感染难以控制或患者临终时，可考虑活性炭敷料吸收异味", AMBER),
        ("③ 环境除味", "通风换气 · 更换床单窗帘 · 沸石/猫砂吸味剂 · 无味除臭喷雾", TEAL_DK),
    ]
    cw = Inches(3.95)
    gap = Inches(0.24)
    x0 = Inches(0.5)
    for i, (h, t, col) in enumerate(steps):
        x = x0 + i * (cw + gap)
        box(s, x, Inches(3.45), cw, Inches(2.05), fill=PANEL, round_=True)
        box(s, x, Inches(3.45), cw, Inches(0.14), fill=col)
        text(s, x + Inches(0.2), Inches(3.72), cw - Inches(0.4), Inches(0.5),
             [[(h, 14, True, INK)]])
        text(s, x + Inches(0.2), Inches(4.22), cw - Inches(0.4), Inches(1.2),
             [[(t, 11.5, False, BODY)]], line_spacing=1.12)
        if i < 2:
            text(s, x + cw + Inches(0.02), Inches(4.1), Inches(0.2), Inches(0.5),
                 [[("→", 18, True, AMBER)]])
    # 不建议
    box(s, Inches(0.5), Inches(5.72), Inches(12.33), Inches(0.6), fill=AMBER_TINT, round_=True)
    text(s, Inches(0.75), Inches(5.82), Inches(11.9), Inches(0.4),
         [[("不建议", 13, True, AMBER_DK),
           ("用带气味精油、空气清新剂、香薰遮盖异味。", 13, False, BODY)]])
    doctor_box(s, Inches(0.5), Inches(6.4), Inches(12.33), Inches(0.62), [
        "抗生素处方应结合培养结果；清创策略由医生与伤口团队共同决策。",
    ])
    set_notes(s, "【讲什么】轻柔清洗、避免激进清创；异味处理三路径（除病因→活性炭→环境除味）；不建议香薰遮盖。\n"
                 "【强调】医生抗生素处方应结合培养结果，清创由医生与伤口团队共同决策。\n"
                 "【追问】活性炭敷料何时使用？异味是否提示感染需升级抗生素？")
    return s


# ============================================================
# P14 症状管理二：出血与渗液
# ============================================================
def page_symptom2():
    s = prs.slides.add_slide(BLANK)
    header(s, "症状管理二：出血与渗液", "模块四 · 医生临床决策要点", 14)
    lx = Inches(0.5)
    rw = Inches(6.0)
    rx = Inches(6.75)
    # 左：出血
    box(s, lx, Inches(1.45), rw, Inches(4.55), fill=PANEL, round_=True)
    box(s, lx, Inches(1.45), rw, Inches(0.14), fill=AMBER)
    text(s, lx + Inches(0.25), Inches(1.7), rw - Inches(0.5), Inches(0.4),
         [[("出血管理", 16, True, INK)]])
    bullet_text(s, lx + Inches(0.25), Inches(2.18), rw - Inches(0.5), Inches(3.7), [
        ("止血：", "创面渗血可按压止血，或使用止血敷料并加压包扎"),
        ("外科介入：", "持续活动性出血难以控制时建议外科介入"),
        ("敷料选择：", "非黏性或可促进止血的壳聚糖、藻酸钙、胶原蛋白、明胶海绵敷料"),
        ("移除敷料：", "动作轻柔，可用生理盐水充分湿润后移除"),
    ], size=13, gap=0.32)
    # 右：渗液
    box(s, rx, Inches(1.45), rw, Inches(4.55), fill=PANEL, round_=True)
    box(s, rx, Inches(1.45), rw, Inches(0.14), fill=TEAL)
    text(s, rx + Inches(0.25), Inches(1.7), rw - Inches(0.5), Inches(0.4),
         [[("渗液管理", 16, True, INK)]])
    bullet_text(s, rx + Inches(0.25), Inches(2.18), rw - Inches(0.5), Inches(3.7), [
        ("高吸收：", "大量渗液优先选择高吸收性敷料"),
        ("难控渗液：", "可联合引流袋或负压伤口治疗（NPWT）"),
        ("保护皮肤：", "必要时使用皮肤保护膜等屏障产品保护周围皮肤"),
    ], size=13, gap=0.32)
    doctor_box(s, Inches(0.5), Inches(6.15), Inches(12.33), Inches(0.8), [
        "掌握外科介入指征；关注出血与凝血风险；负压伤口治疗等方案需要医嘱支持。",
    ])
    set_notes(s, "【讲什么】出血（按压/止血敷料/外科介入/轻柔移除）；渗液（高吸收敷料/引流袋/负压/屏障保护）。\n"
                 "【强调】医生掌握外科介入指征、关注出血与凝血风险，负压治疗需医嘱。\n"
                 "【追问】负压伤口治疗在终末期患者的获益与负担如何权衡？")
    return s


# ============================================================
# P15 症状管理三：疼痛与失禁
# ============================================================
def page_symptom3():
    s = prs.slides.add_slide(BLANK)
    header(s, "症状管理三：疼痛与失禁", "模块四 · 医生临床决策要点", 15)
    lx = Inches(0.5)
    rw = Inches(6.0)
    rx = Inches(6.75)
    # 左：疼痛
    box(s, lx, Inches(1.45), rw, Inches(4.4), fill=PANEL, round_=True)
    box(s, lx, Inches(1.45), rw, Inches(0.14), fill=TEAL)
    text(s, lx + Inches(0.25), Inches(1.7), rw - Inches(0.5), Inches(0.4),
         [[("疼痛管理", 16, True, INK)]])
    bullet_text(s, lx + Inches(0.25), Inches(2.18), rw - Inches(0.5), Inches(3.6), [
        ("区分疼痛：", "全身/伤口因素所致慢性疼痛 vs 翻身、换药、清创等操作所致急性疼痛，按原因给予恰当镇痛"),
        ("癌性疼痛：", "按 WHO 镇痛阶梯方案管理"),
        ("非癌性/操作相关急性疼痛：", "有计划地补充短效止痛药物，并结合音乐、放松等非药物干预"),
    ], size=13, gap=0.34)
    # 右：失禁
    box(s, rx, Inches(1.45), rw, Inches(4.4), fill=PANEL, round_=True)
    box(s, rx, Inches(1.45), rw, Inches(0.14), fill=AMBER)
    text(s, rx + Inches(0.25), Inches(1.7), rw - Inches(0.5), Inches(0.4),
         [[("失禁管理", 16, True, INK)]])
    bullet_text(s, rx + Inches(0.25), Inches(2.18), rw - Inches(0.5), Inches(2.2), [
        "视情况留置尿管或粪便收集装置",
        "失禁后及时清洁皮肤并使用屏障产品保护",
    ], size=13, gap=0.3)
    # 失禁数据
    box(s, rx + Inches(0.25), Inches(3.55), rw - Inches(0.5), Inches(1.9), fill=TEAL_TINT, round_=True)
    text(s, rx + Inches(0.45), Inches(3.7), rw - Inches(0.9), Inches(0.55),
         [[("终末期失禁发生率", 12, True, TEAL_DK)]])
    text(s, rx + Inches(0.45), Inches(4.2), rw - Inches(0.9), Inches(0.5),
         [[("便失禁 约 64.5%", 17, True, TEAL_DK)]])
    text(s, rx + Inches(0.45), Inches(4.75), rw - Inches(0.9), Inches(0.5),
         [[("尿失禁 约 72%~77%", 17, True, TEAL_DK)]])
    doctor_box(s, Inches(0.5), Inches(6.05), Inches(12.33), Inches(0.9), [
        "镇痛处方与阿片类药物的规范使用；导管或粪便收集装置的留置决策。",
    ])
    set_notes(s, "【讲什么】疼痛区分慢性/操作相关急性、癌性疼痛WHO阶梯、非药物干预；失禁管理（便失禁约64.5%、尿失禁约72%~77%）。\n"
                 "【强调】医生关注镇痛处方、阿片类规范使用、导管/粪便收集装置决策。\n"
                 "【追问】阿片类在终末期患者的呼吸抑制顾虑如何权衡？")
    return s


# ============================================================
# P16 健康教育
# ============================================================
def page_education():
    s = prs.slides.add_slide(BLANK)
    header(s, "健康教育", "模块四 · 医生临床决策要点", 16)
    # 左：教育内容清单
    lx = Inches(0.5)
    lw = Inches(6.0)
    text(s, lx, Inches(1.45), lw, Inches(0.35), [[("结构化健康教育内容", 16, True, TEAL_DK)]])
    items = [
        ("护理目标与计划", "让患者/照顾者理解照护方向"),
        ("PI 相关知识", "成因、风险、识别"),
        ("预防和护理策略", "可操作的具体措施"),
    ]
    y = Inches(1.95)
    for i, (h, t) in enumerate(items):
        box(s, lx, y, lw, Inches(1.05), fill=PANEL, round_=True)
        box(s, lx, y, Inches(0.08), Inches(1.05), fill=TEAL)
        text(s, lx + Inches(0.25), y + Inches(0.12), lw - Inches(0.45), Inches(0.4),
             [[(h, 14, True, INK)]])
        text(s, lx + Inches(0.25), y + Inches(0.52), lw - Inches(0.45), Inches(0.4),
             [[(t, 12, False, BODY)]])
        y += Inches(1.2)
    text(s, lx, Inches(5.7), lw, Inches(1.0),
         [[("通过", 13, False, BODY),
           ("共同决策、个体化教育、技能培训与心理支持", 13, True, TEAL_DK),
           ("，促进患者及照顾者参与全程管理。", 13, False, BODY)]], line_spacing=1.2)
    # 右：沟通要点
    rx = Inches(6.75)
    rw = Inches(6.08)
    text(s, rx, Inches(1.45), rw, Inches(0.35), [[("向家属说明的沟通要点", 16, True, TEAL_DK)]])
    box(s, rx, Inches(1.95), rw, Inches(1.7), fill=AMBER_TINT, round_=True)
    box(s, rx, Inches(1.95), Inches(0.08), Inches(1.7), fill=AMBER)
    text(s, rx + Inches(0.28), Inches(2.15), rw - Inches(0.55), Inches(1.35),
         [[("终末期患者的皮肤损伤，本质上是多器官功能衰竭的表现，并非护理不当所致。", 15, True, AMBER_DK)]],
         line_spacing=1.25)
    doctor_box(s, rx, Inches(3.9), rw, Inches(1.6), [
        "参与家属沟通与共同决策；",
        "帮助家属建立合理预期，理解「不可完全避免」的现实。",
    ])
    set_notes(s, "【讲什么】结构化健康教育内容（目标计划/PI知识/预防护理）；通过共同决策、技能培训、心理支持促进全程参与；说明皮肤损伤是多器官功能衰竭表现而非护理不当。\n"
                 "【强调】医生参与沟通与共同决策，帮助家属建立合理预期。\n"
                 "【追问】如何向家属解释「不可避免」又不让家属感觉照护失职？")
    return s


# ============================================================
# P17 转介会诊与质量控制
# ============================================================
def page_referral():
    s = prs.slides.add_slide(BLANK)
    header(s, "转介会诊与质量控制", "模块四 · 医生临床决策要点", 17)
    lx = Inches(0.5)
    lw = Inches(5.9)
    # 左：转介指征
    text(s, lx, Inches(1.42), lw, Inches(0.35), [[("转介 / 会诊指征", 16, True, TEAL_DK)]])
    refs = [
        ("伤口症状恶化", "疼痛加剧、严重出血"),
        ("全身状况", "败血症、意识障碍、心理问题等"),
    ]
    y = Inches(1.92)
    for h, t in refs:
        box(s, lx, y, lw, Inches(1.05), fill=PANEL, round_=True)
        box(s, lx, y, Inches(0.08), Inches(1.05), fill=AMBER)
        text(s, lx + Inches(0.25), y + Inches(0.12), lw - Inches(0.45), Inches(0.4),
             [[(h, 14, True, INK)]])
        text(s, lx + Inches(0.25), y + Inches(0.52), lw - Inches(0.45), Inches(0.4),
             [[(t, 12, False, BODY)]])
        y += Inches(1.18)
    text(s, lx, Inches(4.4), lw, Inches(1.4),
         [[("出现上述情形时，应邀请相关专科介入。", 13, True, BODY)]], line_spacing=1.2)
    doctor_box(s, lx, Inches(5.35), lw, Inches(1.5), [
        "掌握会诊指征；",
        "参与质量指标设计与数据审阅。",
    ])
    # 右：质量环
    rx = Inches(6.7)
    rw = Inches(6.13)
    text(s, rx, Inches(1.42), rw, Inches(0.35), [[("质量控制", 16, True, TEAL_DK)]])
    box(s, rx, Inches(1.82), rw, Inches(1.15), fill=TEAL, round_=True)
    text(s, rx + Inches(0.28), Inches(1.92), rw - Inches(0.55), Inches(0.95),
         [[("质控核心不是「零发生」", 16, True, WHITE)],
          [("而是确保护理过程符合「舒适、尊严」的终末期照护目标", 12.5, False, RGBColor(0xDD, 0xEC, 0xEA))]],
         line_spacing=1.2)
    text(s, rx, Inches(3.15), rw, Inches(0.35), [[("过程指标", 14, True, INK)]])
    box(s, rx, Inches(3.5), rw, Inches(1.05), fill=PANEL, round_=True)
    text(s, rx + Inches(0.25), Inches(3.62), rw - Inches(0.5), Inches(0.85),
         [[("风险评估记录 · 护理计划 · 健康教育知晓率 · 伤口疼痛评分", 12.5, False, BODY)]], line_spacing=1.2)
    text(s, rx, Inches(4.7), rw, Inches(0.35), [[("结局指标", 14, True, INK)]])
    box(s, rx, Inches(5.05), rw, Inches(1.05), fill=AMBER_TINT, round_=True)
    text(s, rx + Inches(0.25), Inches(5.17), rw - Inches(0.5), Inches(0.85),
         [[("满意度 · 舒适度 · PI 现患率 · 2 期及以上发生率", 12.5, False, BODY)]], line_spacing=1.2)
    set_notes(s, "【讲什么】转介/会诊指征（疼痛加剧、严重出血、败血症、意识障碍、心理问题）；过程+结局质量指标体系；质控核心是舒适尊严而非零发生。\n"
                 "【强调】医生掌握会诊指征、参与质量指标设计与数据审阅。\n"
                 "【追问】质量指标如何采集落地？「零发生」为何不现实？")
    return s


# ============================================================
# P18 讨论、证据局限与未来方向
# ============================================================
def page_discussion():
    s = prs.slides.add_slide(BLANK)
    header(s, "讨论、局限与未来方向", "模块五 · 讨论与展望", 18)
    # 上半：讨论与局限
    text(s, Inches(0.5), Inches(1.42), Inches(6.0), Inches(0.35),
         [[("讨论与证据局限", 16, True, TEAL_DK)]])
    box(s, Inches(0.5), Inches(1.82), Inches(12.33), Inches(2.0), fill=PANEL, round_=True)
    bullet_text(s, Inches(0.8), Inches(2.0), Inches(11.8), Inches(1.7), [
        ("弥补空白：", "本共识弥补了现有指南对终末期患者关注不足的问题。"),
        ("证据局限：", "部分推荐意见证据等级较低，多来源于专家经验与低质量研究，临床应用中应结合患者情况灵活调整。"),
        ("核心挑战：", "在「积极预防」与「避免过度干预」之间取得平衡。"),
    ], size=13.5, gap=0.32)
    # 下半：未来四方向
    text(s, Inches(0.5), Inches(4.05), Inches(6.0), Inches(0.35),
         [[("未来研究方向（四方向）", 16, True, TEAL_DK)]])
    dirs = [
        ("工具开发", "终末期特异性风险评估工具，多中心前瞻验证"),
        ("高质量证据", "高质量 RCT + 卫生经济学 + 患者/家属报告结局"),
        ("质性研究", "探索认知与决策行为"),
        ("实施评价", "发布后多中心前后对照或中断时间序列评价实施效果"),
    ]
    cw = Inches(2.95)
    gap = Inches(0.18)
    x0 = Inches(0.5)
    for i, (h, t) in enumerate(dirs):
        x = x0 + i * (cw + gap)
        box(s, x, Inches(4.5), cw, Inches(2.2), fill=WHITE, line=LINE, line_w=1.0, round_=True)
        box(s, x, Inches(4.5), cw, Inches(0.14), fill=AMBER if i % 2 else TEAL)
        text(s, x + Inches(0.2), Inches(4.76), cw - Inches(0.4), Inches(0.4),
             [[(h, 14.5, True, INK)]])
        text(s, x + Inches(0.2), Inches(5.22), cw - Inches(0.4), Inches(1.4),
             [[(t, 11.5, False, BODY)]], line_spacing=1.15)
    set_notes(s, "【讲什么】共识弥补空白；证据等级总体偏低（专家经验/低质量研究）；核心挑战是「积极预防 vs 避免过度干预」的平衡；未来四方向。\n"
                 "【强调】临床应用需结合患者情况灵活调整，不能机械套用。\n"
                 "【追问】低证据等级的推荐如何在临床安全使用？四个方向中哪个最紧迫？")
    return s


# ============================================================
# P19 总结与临床行动要点
# ============================================================
def page_summary():
    s = prs.slides.add_slide(BLANK)
    header(s, "总结与临床行动要点", "模块五 · 讨论与展望", 19)
    # 三句话总结
    text(s, Inches(0.5), Inches(1.45), Inches(6.0), Inches(0.35),
         [[("三句话总结", 16, True, TEAL_DK)]])
    sums = [
        "共识的核心，是把目标从「愈合」转向「舒适与尊严」。",
        "29 条推荐按 9 大主题，构成可落地的行动框架。",
        "证据存在局限，临床需个体化决策。",
    ]
    y = Inches(1.9)
    for i, t in enumerate(sums, 1):
        box(s, Inches(0.5), y, Inches(12.33), Inches(0.72), fill=PANEL, round_=True)
        box(s, Inches(0.5), y, Inches(0.08), Inches(0.72), fill=TEAL)
        text(s, Inches(0.8), y, Inches(11.9), Inches(0.72), [
            [(f"{i}. ", 14, True, TEAL_DK), (t, 14, False, BODY)],
        ], anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.84)
    # 医生行动清单
    text(s, Inches(0.5), Inches(4.55), Inches(6.0), Inches(0.35),
         [[("医生行动清单", 16, True, TEAL_DK)]])
    acts = [
        ("复核", "风险评估与高危药物"),
        ("参与", "多学科团队与家属沟通"),
        ("掌握", "会诊与转介指征"),
        ("参与", "质量指标数据审阅"),
    ]
    cw = Inches(2.95)
    gap = Inches(0.18)
    x0 = Inches(0.5)
    for i, (h, t) in enumerate(acts):
        x = x0 + i * (cw + gap)
        box(s, x, Inches(5.05), cw, Inches(1.7), fill=TEAL_TINT, round_=True)
        box(s, x, Inches(5.05), cw, Inches(0.14), fill=AMBER)
        text(s, x + Inches(0.2), Inches(5.3), cw - Inches(0.4), Inches(0.55),
             [[(f"0{i+1}  {h}", 16, True, TEAL_DK)]])
        text(s, x + Inches(0.2), Inches(5.95), cw - Inches(0.4), Inches(0.7),
             [[(t, 12.5, False, BODY)]], line_spacing=1.1)
    set_notes(s, "【讲什么】三句话总结（愈合→舒适尊严；29条×9主题行动框架；证据局限需个体化）+ 医生四项行动清单。\n"
                 "【强调】把共识落到可执行动作：复核评估、参与MDT与沟通、掌握转介、参与质控。\n"
                 "【追问】本院如何将这四项行动落地为日常流程？")
    return s


# ============================================================
# P20 感谢聆听
# ============================================================
def page_closing():
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, SLIDE_W, Inches(0.10), fill=TEAL)
    box(s, 0, Inches(7.4), SLIDE_W, Inches(0.10), fill=TEAL)
    text(s, Inches(0.5), Inches(2.55), Inches(12.33), Inches(1.0),
         [[("感谢聆听", 44, True, INK)]], align=PP_ALIGN.CENTER)
    text(s, Inches(0.5), Inches(3.65), Inches(12.33), Inches(0.5),
         [[("欢迎交流讨论", 20, False, BODY)]], align=PP_ALIGN.CENTER)
    box(s, Inches(5.42), Inches(4.35), Inches(2.5), Pt(2.2), fill=AMBER)
    text(s, Inches(1.0), Inches(4.75), Inches(11.33), Inches(1.6), [
        [("来源：", 12, True, MUTED), ("中华老年医学杂志 2026年6月 第45卷第6期", 12, False, MUTED)],
        [("DOI：", 12, True, MUTED), ("10.3760/cma.j.issn.0254-9026.2026.06.003", 12, False, MUTED)],
        [("基金：", 12, True, MUTED), ("国家重点研发计划（2023YFC3605900）", 12, False, MUTED)],
        [("作者声明无利益冲突", 12, False, MUTED)],
    ], align=PP_ALIGN.CENTER, line_spacing=1.25)
    badge = box(s, Inches(5.72), Inches(6.15), Inches(1.9), Inches(0.46), fill=TEAL_TINT,
                line=TEAL, line_w=1.0, round_=True)
    shape_text(badge, "内部教学使用", 13, True, TEAL_DK)
    footer(s, 20)
    set_notes(s, "【讲什么】致谢并引导讨论；底部展示来源、DOI、基金与利益冲突声明。\n"
                 "【强调】标注内部教学使用、来源与版权合规。\n"
                 "【追问】（无）")
    return s


# ============================================================
def main():
    pages = [
        page_cover, page_toc, page_background, page_method1, page_method2,
        page_method3, page_definition, page_overview, page_mdt, page_assessment,
        page_prevention1, page_prevention2, page_symptom1, page_symptom2,
        page_symptom3, page_education, page_referral, page_discussion,
        page_summary, page_closing,
    ]
    for fn in pages:
        fn()
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'output', '终末期患者压力性损伤护理专家共识2026_内部教学.pptx')
    prs.save(out)
    print("已生成:", out, "共", len(prs.slides.__iter__.__self__._sldIdLst), "页")


if __name__ == "__main__":
    main()
