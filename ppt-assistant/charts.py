"""Validated native charts with embedded workbooks, shared styling and CSV input."""
from __future__ import annotations

import copy
import csv
import math
from pathlib import Path

from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_DATA_LABEL_POSITION,
                             XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK)
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

KINDS = {
    'column': XL_CHART_TYPE.COLUMN_CLUSTERED,
    'bar': XL_CHART_TYPE.BAR_CLUSTERED,
    'line': XL_CHART_TYPE.LINE_MARKERS,
    'area': XL_CHART_TYPE.AREA,
    'stacked_column': XL_CHART_TYPE.COLUMN_STACKED,
    'stacked_bar': XL_CHART_TYPE.BAR_STACKED,
    'percent_column': XL_CHART_TYPE.COLUMN_STACKED_100,
    'percent_bar': XL_CHART_TYPE.BAR_STACKED_100,
    'donut': XL_CHART_TYPE.DOUGHNUT,
    'pie': XL_CHART_TYPE.PIE,
    'scatter': XL_CHART_TYPE.XY_SCATTER,
    'bubble': XL_CHART_TYPE.BUBBLE,
}
PALETTE = ['176C88', 'D97706', '6366F1', '2A9D8F', 'C0392B', '64748B']


def _numeric(value):
    return type(value) in {int, float} and math.isfinite(value)


def _color(value):
    return (isinstance(value, str) and len(value) == 6
            and all(c in '0123456789abcdefABCDEF' for c in value))


def resolve_chart(spec: dict, base: Path) -> dict:
    """Copy, load optional CSV and validate; never infer or synthesize missing values."""
    if not isinstance(spec, dict):
        raise ValueError('chart must be an object')
    chart = copy.deepcopy(spec)
    kind = chart.setdefault('kind', 'column')
    if not isinstance(kind, str) or kind not in KINDS:
        raise ValueError(f'unsupported chart kind: {kind}')
    series = chart.get('series')
    if not isinstance(series, list) or not 1 <= len(series) <= 6:
        raise ValueError('chart requires 1–6 series')
    if any(not isinstance(s, dict) or not isinstance(s.get('name'), str) or not s['name'].strip()
           for s in series):
        raise ValueError('every chart series needs a name')
    if 'number_format' in chart and (not isinstance(chart['number_format'], str)
                                       or not chart['number_format']):
        raise ValueError('number_format must be a non-empty Excel format string')
    if 'data_labels' in chart and type(chart['data_labels']) is not bool:
        raise ValueError('data_labels must be boolean')
    if 'data_file' in chart:
        if kind in {'scatter', 'bubble'}:
            raise ValueError('XY charts use explicit points, not category CSV')
        if 'categories' in chart or any('values' in s for s in chart.get('series', [])):
            raise ValueError('CSV and inline chart values cannot be mixed')
        if not isinstance(chart['data_file'], str) or not chart['data_file'].strip():
            raise ValueError('data_file must be a CSV path')
        path = base / chart['data_file']
        if path.suffix.lower() != '.csv':
            raise ValueError('chart data_file must be a CSV file')
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            rows = list(reader)
            columns = [chart.get('category_column')] + [s.get('column') for s in chart.get('series', [])]
            if any(not isinstance(c, str) or not c or c not in (reader.fieldnames or []) for c in columns):
                raise ValueError('CSV category/series column is missing')
            if len(set(reader.fieldnames)) != len(reader.fieldnames) or any(
                    None in row or any(value is None for value in row.values()) for row in rows):
                raise ValueError('CSV requires unique headers and complete rows; use empty cells for gaps')
        chart['categories'] = [row[chart['category_column']] for row in rows]
        for item in series:
            try:
                item['values'] = [float(row[item['column']]) if row[item['column']].strip()
                                    else None for row in rows]
            except (ValueError, AttributeError, KeyError) as error:
                raise ValueError(f"CSV series {item.get('name')} contains invalid cells") from error
    if kind in {'scatter', 'bubble'}:
        count = 3 if kind == 'bubble' else 2
        for s in series:
            points = s.get('points')
            if not isinstance(points, list) or not 1 <= len(points) <= 500:
                raise ValueError('XY series requires 1–500 points')
            if any(not isinstance(p, list) or len(p) != count or not all(_numeric(v) for v in p)
                   or (count == 3 and p[2] <= 0) for p in points):
                raise ValueError('XY points must be finite [x,y] or [x,y,positive_size]')
    else:
        categories = chart.get('categories')
        if (not isinstance(categories, list) or not 1 <= len(categories) <= 40
                or any(not isinstance(c, str) or not c.strip() for c in categories)):
            raise ValueError('chart requires 1–40 non-empty category labels')
        for s in series:
            values = s.get('values')
            if (not isinstance(values, list) or len(values) != len(categories)
                    or any(v is not None and not _numeric(v) for v in values)
                    or all(v is None for v in values)):
                raise ValueError('chart values must match categories and contain finite numbers or null gaps')
        if kind in {'pie', 'donut'}:
            values = series[0]['values']
            if (len(series) != 1 or len(categories) > 8
                    or any(v is None or v < 0 for v in values) or sum(values) <= 0):
                raise ValueError('pie/donut requires one non-negative series, positive total and at most 8 slices')
        if kind.startswith('percent_'):
            if '%' in chart.get('number_format', ''):
                raise ValueError('percent stacks retain source counts; use y_axis.number_format for percentages')
            if any(v is None or v < 0 for s in series for v in s['values']):
                raise ValueError('percent stacks require non-negative, complete data')
            if any(sum(s['values'][i] for s in series) <= 0 for i in range(len(categories))):
                raise ValueError('each percent stack needs a positive total')
        if chart.get('sort'):
            if kind not in {'bar', 'column'} or len(series) != 1 or chart['sort'] not in ('ascending', 'descending'):
                raise ValueError('sorting requires a single bar/column series and ascending/descending')
            order = sorted(range(len(categories)), key=lambda i: (
                series[0]['values'][i] is None,
                (-series[0]['values'][i] if chart['sort'] == 'descending' else series[0]['values'][i])
                if series[0]['values'][i] is not None else 0))
            chart['categories'] = [categories[i] for i in order]
            series[0]['values'] = [series[0]['values'][i] for i in order]
    for s in series:
        if 'color' in s and not _color(s['color']):
            raise ValueError('series.color must be six hexadecimal digits')
    for field in ('highlight',):
        if not isinstance(chart.get(field, []), list) or any(
                c not in chart.get('categories', []) for c in chart.get(field, [])):
            raise ValueError('highlight must contain existing category labels')
    for axis_name in ('x_axis', 'y_axis'):
        axis = chart.get(axis_name, {})
        if not isinstance(axis, dict):
            raise ValueError(f'{axis_name} must be an object')
        for field in ('title', 'number_format'):
            if field in axis and not isinstance(axis[field], str):
                raise ValueError(f'{axis_name}.{field} must be a string')
        if 'font_size' in axis and (not _numeric(axis['font_size']) or not 9 <= axis['font_size'] <= 18):
            raise ValueError('axis font_size must be between 9 and 18 points')
        if 'label_interval' in axis and (axis_name != 'x_axis' or kind in {'scatter', 'bubble', 'pie', 'donut'}
                                        or type(axis['label_interval']) is not int
                                        or not 1 <= axis['label_interval'] <= 40):
            raise ValueError('label_interval requires a category axis and an integer from 1 to 40')
        for field in ('min', 'max', 'major_unit'):
            if field in axis and not _numeric(axis[field]):
                raise ValueError(f'{axis_name}.{field} must be finite')
        if 'major_unit' in axis and axis['major_unit'] <= 0:
            raise ValueError('axis major_unit must be positive')
        if kind.startswith('percent_') and axis_name == 'y_axis' and (
                axis.get('min', 0) != 0 or axis.get('max', 1) != 1
                or axis.get('major_unit', .2) > 1):
            raise ValueError('percent stack value axis uses 0–1, displayed as 0–100%')
        if 'min' in axis and 'max' in axis and axis['min'] >= axis['max']:
            raise ValueError('axis minimum must be below maximum')
        if kind not in {'scatter', 'bubble', 'pie', 'donut'} and axis_name == 'x_axis' and any(
                field in axis for field in ('min', 'max', 'major_unit')):
            raise ValueError('category axis has no numeric bounds; use y_axis for value bounds, including bar charts')
        if kind not in {'pie', 'donut'} and (axis_name == 'y_axis' or kind in {'scatter', 'bubble'}):
            if kind.startswith('percent_'):
                values = [0., 1.]
            elif kind in {'scatter', 'bubble'}:
                dimension = 0 if axis_name == 'x_axis' else 1
                values = [p[dimension] for s in series for p in s['points']]
            else:
                values = [v for s in series for v in s['values'] if v is not None]
                if 'stacked' in kind:
                    values += [sum(max(s['values'][i] or 0, 0) for s in series)
                               for i in range(len(chart['categories']))]
                    values += [sum(min(s['values'][i] or 0, 0) for s in series)
                               for i in range(len(chart['categories']))]
            if any(v < axis.get('min', -math.inf) or v > axis.get('max', math.inf) for v in values):
                raise ValueError('axis bounds would hide source values')
            if kind in {'column', 'bar', 'area', 'stacked_column', 'stacked_bar'} and (
                    axis.get('min', 0) > 0 or axis.get('max', 0) < 0):
                raise ValueError('bar and area value axes must include zero')
    return chart


def draw_chart(slide, spec: dict, base: Path, box: tuple, font: str, colors: dict):
    """Add a native PowerPoint chart and its editable workbook; box uses inches."""
    spec = resolve_chart(spec, base)
    kind = spec['kind']
    fmt = spec.get('number_format', '0.##')
    if kind in {'scatter', 'bubble'}:
        data = BubbleChartData() if kind == 'bubble' else XyChartData()
        for s in spec['series']:
            target = data.add_series(s['name'])
            for point in s['points']:
                target.add_data_point(*point)
    else:
        data = CategoryChartData()
        data.categories = spec['categories']
        for s in spec['series']:
            data.add_series(s['name'], s['values'], fmt)
    x, y, w, h = box
    if kind in {'bar', 'stacked_bar', 'percent_bar'} and len(spec['categories']) > max(6, int(h * 3)):
        raise ValueError('too many bar categories for the panel height; split the chart across pages')
    chart = slide.shapes.add_chart(KINDS[kind], Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    # Explicit chart-area fill avoids theme white backgrounds in dark decks.
    properties = OxmlElement('c:spPr')
    fill = OxmlElement('a:solidFill')
    rgb = OxmlElement('a:srgbClr')
    rgb.set('val', colors['panel'])
    fill.append(rgb)
    properties.append(fill)
    line = OxmlElement('a:ln')
    line.append(OxmlElement('a:noFill'))
    properties.append(line)
    properties.append(OxmlElement('a:effectLst'))
    chart._chartSpace.insert_element_before(properties, 'c:txPr', 'c:externalData',
                                           'c:printSettings', 'c:userShapes', 'c:extLst')
    chart.has_title = False
    chart.has_legend = len(spec['series']) > 1 or kind in {'pie', 'donut'}
    compact_composition = kind in {'pie', 'donut'} and h < 2.6
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.RIGHT if compact_composition else XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = font
        chart.legend.font.size = Pt(11 if w < 7 else 13)
        chart.legend.font.color.rgb = RGBColor.from_string(colors['text'])
    plot = chart.plots[0]
    if kind in {'pie', 'donut'}:
        area = chart._chartSpace.xpath('c:chart/c:plotArea')[0]
        layout = OxmlElement('c:layout')
        area.insert(0, layout)
        manual = OxmlElement('c:manualLayout')
        for tag, value in [('layoutTarget', 'inner'), ('xMode', 'factor'), ('yMode', 'factor'),
                           ('wMode', 'factor'), ('hMode', 'factor'),
                           ('x', '.05'), ('y', '.02'),
                           ('w', '.7' if compact_composition else '.9'),
                           ('h', '.95' if compact_composition else '.78')]:
            element = OxmlElement('c:' + tag)
            element.set('val', value)
            manual.append(element)
        layout.append(manual)
    if kind == 'donut':
        plot.hole_size = 65
    if kind not in {'scatter', 'bubble', 'area'}:
        plot.has_data_labels = spec.get('data_labels', not kind.startswith('percent_')
                                       and len(spec.get('categories', [])) <= 8)
        if plot.has_data_labels:
            labels = plot.data_labels
            labels.font.name, labels.font.size = font, Pt(10 if w < 7 else 13)
            labels.font.color.rgb = RGBColor.from_string(colors['text'])
            if kind in {'pie', 'donut'}:
                labels.show_value, labels.show_percentage = False, True
                labels.position = XL_DATA_LABEL_POSITION.CENTER
                labels.font.color.rgb = RGBColor.from_string('FFFFFF')
            else:
                labels.number_format = fmt
                labels.show_value = True
                labels.position = (XL_DATA_LABEL_POSITION.CENTER if 'stacked' in kind or kind.startswith('percent_')
                                   else XL_DATA_LABEL_POSITION.OUTSIDE_END if kind in {'bar', 'column'}
                                   else XL_DATA_LABEL_POSITION.ABOVE)
                if 'stacked' in kind or kind.startswith('percent_'):
                    labels.font.color.rgb = RGBColor.from_string('FFFFFF')
    for i, series in enumerate(chart.series):
        color = spec['series'][i].get('color', PALETTE[i])
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
        series.format.line.color.rgb = RGBColor.from_string(color)
        if kind in {'line', 'scatter'}:
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = 6
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
            series.marker.format.line.color.rgb = RGBColor.from_string(color)
            series.format.line.width = Pt(2.2)
            if kind == 'scatter':
                series.format.line.fill.background()
        if kind in {'pie', 'donut'} or (len(chart.series) == 1 and spec.get('highlight')):
            for j, point in enumerate(series.points):
                point_color = PALETTE[j % len(PALETTE)] if kind in {'pie', 'donut'} else (
                    colors['accent'] if spec['categories'][j] in spec['highlight'] else '94A3B8')
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(point_color)
                point.format.line.fill.background()
    if kind not in {'pie', 'donut'}:
        for name, axis in [('x_axis', chart.category_axis), ('y_axis', chart.value_axis)]:
            options = spec.get(name, {})
            axis.tick_labels.font.name = font
            axis.tick_labels.font.size = Pt(options.get('font_size', 10 if w < 7 else 12))
            axis.tick_labels.font.color.rgb = RGBColor.from_string(colors['muted'])
            axis.format.line.color.rgb = RGBColor.from_string(colors['muted'])
            axis.major_tick_mark = XL_TICK_MARK.NONE
            axis.minor_tick_mark = XL_TICK_MARK.NONE
            axis.has_title = bool(options.get('title'))
            if axis.has_title:
                frame = axis.axis_title.text_frame
                frame.text = options['title']
                for p in frame.paragraphs:
                    p.font.name, p.font.size = font, Pt(11)
                    p.font.color.rgb = RGBColor.from_string(colors['muted'])
            if name == 'y_axis' or kind in {'scatter', 'bubble'}:
                axis.tick_labels.number_format = options.get('number_format', '0%' if kind.startswith('percent_') else fmt)
                for field, attribute in [('min', 'minimum_scale'), ('max', 'maximum_scale'), ('major_unit', 'major_unit')]:
                    if field in options:
                        setattr(axis, attribute, options[field])
                if name == 'y_axis' and kind not in {'scatter', 'bubble', 'line'} and 'min' not in options:
                    if all(v is None or v >= 0 for s in spec['series'] for v in s['values']):
                        axis.minimum_scale = 0
                axis.has_major_gridlines = name == 'y_axis'
                if axis.has_major_gridlines:
                    axis.major_gridlines.format.line.color.rgb = RGBColor.from_string('D5DCE5' if colors['panel'] == 'FFFFFF' else '3D506B')
                    axis.major_gridlines.format.line.width = Pt(.5)
            else:
                interval = options.get('label_interval', max(1, math.ceil(len(spec['categories']) / max(4, int(w * 1.2)))))
                skip = OxmlElement('c:tickLblSkip')
                skip.set('val', str(interval))
                axis._element.insert_element_before(skip, 'c:tickMarkSkip', 'c:noMultiLvlLbl', 'c:extLst')
        if kind in {'bar', 'stacked_bar', 'percent_bar'}:
            chart.category_axis.reverse_order = True
            chart.value_axis._element.xpath('c:crosses')[0].set('val', 'max')
    return chart
