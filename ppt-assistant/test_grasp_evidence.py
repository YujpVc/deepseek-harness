"""Checks for paired depth colors, physical projection, and text obstruction."""
import json
from pathlib import Path
import tempfile
import unittest

import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches, Pt

from grasp_evidence import (candidate_projections, depth_pair,
                            gated_candidate_projections, gripper_projection)
from check import check


class GraspEvidenceTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)

    def test_equal_depth_has_equal_color_and_invalid_pixels_are_gray(self):
        np.save(self.base / "before.npy", [[.95, 0], [1.0, float("nan")]])
        np.save(self.base / "after.npy", [[.95, 1.1], [.9, .99]])
        before, after = depth_pair(self.base / "before.npy", self.base / "after.npy",
                                   [0, 0, 2, 2], [.9, 1.1])
        self.assertEqual(before.getpixel((0, 0)), after.getpixel((0, 0)))
        self.assertEqual(before.getpixel((1, 0)), (190, 190, 190))
        self.assertEqual(before.getpixel((1, 1)), (190, 190, 190))

    def test_mismatched_depth_array_shapes_are_rejected(self):
        np.save(self.base / "before.npy", np.ones((2, 2)))
        np.save(self.base / "after.npy", np.ones((3, 3)))
        with self.assertRaisesRegex(ValueError, "matching"):
            depth_pair(self.base / "before.npy", self.base / "after.npy", [0, 0, 2, 2], [.9, 1.1])

    def test_recorded_opening_projects_using_camera_intrinsics(self):
        transform = np.eye(4)
        transform[2, 3] = 1
        (self.base / "pose.json").write_text(json.dumps({
            "T_camera_cgn_gripper": transform.tolist(), "gripper_opening_m": .02}))
        (self.base / "camera.json").write_text(json.dumps({
            "color_k": [[100, 0, 50], [0, 100, 50], [0, 0, 1]]}))
        result = gripper_projection(self.base / "pose.json", self.base / "camera.json")
        self.assertEqual(result["opening_mm"], 20)
        self.assertEqual(result["center_uv"], [50, 50])
        # Gap at the tips must reflect the opening and perspective, rather
        # than a hand-drawn width selected for appearance.
        left_tip = result["segments"][0][1]
        right_tip = result["segments"][2][1]
        self.assertAlmostEqual(right_tip[0] - left_tip[0], 100 * .02 / 1.1034)

    def test_crossing_wireframes_do_not_hide_real_text_overlap(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for coords in [(1, 1, 3, 3), (1, 3, 3, 1)]:
            slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *[Inches(v) for v in coords])
        path = self.base / "wireframe.pptx"
        prs.save(path)
        self.assertEqual(check(str(path))["overlap"], [])
        for label in ["First", "Second"]:
            shape = slide.shapes.add_textbox(Inches(5), Inches(1), Inches(2), Inches(1))
            shape.text = label
            shape.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        prs.save(path)
        self.assertEqual(len(check(str(path))["overlap"]), 1)

    def test_candidates_rank_by_score_with_the_selected_pose_geometry(self):
        transform = np.eye(4)
        transform[2, 3] = 1
        camera = {"color_k": [[100, 0, 50], [0, 100, 50], [0, 0, 1]]}
        (self.base / "camera.json").write_text(json.dumps(camera))
        record = {"T_camera_cgn_gripper": transform.tolist(), "gripper_opening_m": .06}
        (self.base / "pose.json").write_text(json.dumps(record))
        response = {"grasps": {"total_grasps": 2,
            "pred_grasps_cam": {"0": [transform.tolist(), transform.tolist()]},
            "scores": {"0": [.1, .8]}, "gripper_openings": {"0": [.02, .06]}}}
        path = self.base / "response.json"
        path.write_text(json.dumps(response))
        result = candidate_projections(path, self.base / "camera.json", count=1)
        self.assertEqual(result["total"], 2)
        selected = result["projections"][0]
        self.assertEqual(selected["index"], 1)
        reference = gripper_projection(self.base / "pose.json", self.base / "camera.json")
        self.assertEqual(selected["segments"], reference["segments"])
        self.assertEqual(selected["opening_mm"], reference["opening_mm"])
        response["grasps"]["total_grasps"] = 3
        path.write_text(json.dumps(response))
        with self.assertRaisesRegex(ValueError, "recorded nonzero total"):
            candidate_projections(path, self.base / "camera.json")
        response["grasps"]["gripper_openings"]["0"] = [.02]
        path.write_text(json.dumps(response))
        with self.assertRaisesRegex(ValueError, "matching lengths"):
            candidate_projections(path, self.base / "camera.json")

    def test_gated_candidates_preserve_every_candidate_status(self):
        transform = np.eye(4)
        transform[2, 3] = 1
        (self.base / "camera.json").write_text(json.dumps({
            "color_k": [[100, 0, 50], [0, 100, 50], [0, 0, 1]]}))
        def item(index, opening=.04):
            return {"candidate_index": index, "candidate_score": .1 + index / 100,
                    "T_camera_cgn_gripper": transform.tolist(),
                    "gripper_opening_m": opening}
        record = {"candidates_evaluated": 3,
                  "candidate_acceptances": [item(0)],
                  "candidate_rejections": [item(1), item(2)]}
        path = self.base / "gate.json"
        path.write_text(json.dumps(record))
        result = gated_candidate_projections(path, self.base / "camera.json")
        self.assertEqual((result["total"], result["accepted"], result["rejected"]), (3, 1, 2))
        self.assertEqual([p["status"] for p in result["projections"]],
                         ["accepted", "rejected", "rejected"])
        record["candidate_rejections"].pop()
        path.write_text(json.dumps(record))
        with self.assertRaisesRegex(ValueError, "every evaluated candidate"):
            gated_candidate_projections(path, self.base / "camera.json")


if __name__ == "__main__":
    unittest.main()
