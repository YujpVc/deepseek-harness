#!/usr/bin/env python3
"""Inspect shape ids and patch selected slide geometry/text frames only."""
from __future__ import annotations

import argparse
import copy
import json
import math
import tempfile
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from speaker_notes import dump, slide_parts

ANCHORS = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
ALIGNMENTS = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
FRAME_KEYS = {"word_wrap", "vertical_anchor", "margin_inches", "line_spacing",
              "space_before_pt", "space_after_pt", "alignment"}


def inspect(pptx: Path) -> dict:
    """List top-level shape ids and inch-based boxes in presentation order."""
    prs = Presentation(pptx)
    return {"size_inches": [prs.slide_width / 914400, prs.slide_height / 914400],
            "pages": [{"page": page, "shapes": [
                {"shape_id": shape.shape_id, "name": shape.name,
                 "type": str(shape.shape_type), "text": shape.text if shape.has_text_frame else None,
                 "box_inches": [value / 914400 for value in
                                (shape.left, shape.top, shape.width, shape.height)]}
                for shape in slide.shapes]} for page, slide in enumerate(prs.slides, 1)]}


def number(value, label: str, *, positive: bool = False) -> float:
    if (type(value) not in (int, float) or not math.isfinite(value)
            or value < 0 or (positive and value == 0)):
        raise ValueError(f"{label} must be a finite {'positive' if positive else 'nonnegative'} number")
    return value


def apply_frame(shape, spec: dict) -> None:
    """Repair spacing without replacing paragraphs, runs, fonts or text."""
    if not isinstance(spec, dict) or not spec or spec.keys() - FRAME_KEYS:
        raise ValueError("text_frame must contain supported, explicit settings")
    if not shape.has_text_frame:
        raise ValueError("selected shape has no text frame")
    frame = shape.text_frame
    frame.auto_size = MSO_AUTO_SIZE.NONE
    if "word_wrap" in spec:
        if type(spec["word_wrap"]) is not bool:
            raise ValueError("word_wrap must be boolean")
        frame.word_wrap = spec["word_wrap"]
    if "vertical_anchor" in spec:
        if spec["vertical_anchor"] not in ANCHORS:
            raise ValueError("vertical_anchor must be top, middle or bottom")
        frame.vertical_anchor = ANCHORS[spec["vertical_anchor"]]
    if "margin_inches" in spec:
        margins = spec["margin_inches"]
        if not isinstance(margins, list) or len(margins) != 4:
            raise ValueError("margin_inches must be [left, top, right, bottom]")
        values = [Inches(number(value, "margin")) for value in margins]
        if values[0] + values[2] >= shape.width or values[1] + values[3] >= shape.height:
            raise ValueError("margins must leave positive text space")
        for key, value in zip(("margin_left", "margin_top", "margin_right", "margin_bottom"), values):
            setattr(frame, key, value)
    for paragraph in frame.paragraphs:
        if "line_spacing" in spec:
            paragraph.line_spacing = float(number(spec["line_spacing"], "line_spacing", positive=True))
        for key in ("space_before_pt", "space_after_pt"):
            if key in spec:
                setattr(paragraph, key.removesuffix("_pt"), Pt(number(spec[key], key)))
        if "alignment" in spec:
            if spec["alignment"] not in ALIGNMENTS:
                raise ValueError("alignment must be left, center or right")
            paragraph.alignment = ALIGNMENTS[spec["alignment"]]


def patch(pptx: Path, specification: Path, output: Path) -> dict:
    """Atomically replace selected slide XML; retain all other ZIP payloads."""
    if pptx.resolve() == output.resolve() or specification.resolve() == output.resolve():
        raise ValueError("output must be separate from the inputs")
    operations = json.loads(specification.read_text(encoding="utf-8"))
    if not isinstance(operations, list) or not operations:
        raise ValueError("specification must be a nonempty list")
    with ZipFile(pptx) as archive:
        infos = archive.infolist()
        if len(infos) != len({item.filename for item in infos}) or archive.testzip():
            raise ValueError("invalid or duplicate PPT ZIP entries")
        before = {item.filename: archive.read(item) for item in infos}
    parts = slide_parts(before)
    prs = Presentation(pptx)
    selected, seen = set(), set()
    for operation in operations:
        if (not isinstance(operation, dict) or operation.keys() - {"page", "shape_id", "box_inches", "text_frame"}
                or not ({"box_inches", "text_frame"} & operation.keys())):
            raise ValueError("operation requires page, shape_id and box_inches or text_frame")
        page, shape_id = operation.get("page"), operation.get("shape_id")
        if type(page) is not int or not 1 <= page <= len(prs.slides) or type(shape_id) is not int:
            raise ValueError("invalid page or shape_id")
        if (page, shape_id) in seen:
            raise ValueError("duplicate page/shape_id operation")
        seen.add((page, shape_id))
        matches = [shape for shape in prs.slides[page - 1].shapes if shape.shape_id == shape_id]
        if len(matches) != 1:
            raise ValueError(f"P{page}: shape_id {shape_id} is not a unique top-level shape")
        shape = matches[0]
        if "box_inches" in operation:
            box = operation["box_inches"]
            if not isinstance(box, list) or len(box) != 4:
                raise ValueError("box_inches must be [left, top, width, height]")
            values = [Inches(number(value, "box", positive=index >= 2)) for index, value in enumerate(box)]
            if values[0] + values[2] > prs.slide_width or values[1] + values[3] > prs.slide_height:
                raise ValueError("box is outside slide bounds")
            for key, value in zip(("left", "top", "width", "height"), values):
                setattr(shape, key, value)
        if "text_frame" in operation:
            apply_frame(shape, operation["text_frame"])
        selected.add(page)
    updated = {parts[page - 1]: dump(prs.slides[page - 1]._element) for page in selected}
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix=".slide-patch-") as directory:
        temporary = Path(directory) / "presentation.pptx"
        with ZipFile(temporary, "w") as archive:
            for info in infos:
                archive.writestr(copy.copy(info), updated.get(info.filename, before[info.filename]))
        with ZipFile(temporary) as archive:
            if archive.testzip() or any(archive.read(name) != data for name, data in before.items() if name not in updated):
                raise ValueError("patch changed an unselected package part")
        Presentation(temporary)
        temporary.replace(output)
    return {"patched_pages": sorted(selected), "patched_shapes": len(seen), "other_parts_preserved": True}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    reader = commands.add_parser("inspect")
    reader.add_argument("pptx", type=Path)
    writer = commands.add_parser("apply")
    writer.add_argument("pptx", type=Path)
    writer.add_argument("specification", type=Path)
    writer.add_argument("output", type=Path)
    args = parser.parse_args()
    result = inspect(args.pptx) if args.command == "inspect" else patch(args.pptx, args.specification, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
