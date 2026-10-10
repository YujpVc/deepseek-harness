#!/usr/bin/env python3
"""Typeset MathText formulas as native editable PowerPoint vector shapes."""
from __future__ import annotations

import argparse
import json
import tempfile
from pathlib import Path

from lxml import etree as ET
from matplotlib.mathtext import MathTextParser
from matplotlib.path import Path as MathPath
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def element(parent, tag, **attrs):
    return ET.SubElement(parent, f"{{{A}}}{tag}", **attrs)


def add_formula(slide, expression: str, left: int, top: int, width: int, height: int,
                color: str = "202020", name: str = "Formula"):
    """Fit one supported MathText expression to a box; emit native cubic curves."""
    if width <= 0 or height <= 0:
        raise ValueError("positive formula box required")
    parsed = MathTextParser("path").parse(expression, dpi=72)
    if not parsed.glyphs or parsed.width <= 0 or parsed.height <= 0:
        raise ValueError("formula must contain typeset glyphs")
    scale = min(width / parsed.width, height / parsed.height)
    w, h = int(parsed.width * scale), int(parsed.height * scale)
    shape = slide.shapes.build_freeform(0, 0).add_line_segments([(1, 0), (1, 1)], close=True).convert_to_shape()
    shape.left, shape.top = left, top
    shape.width, shape.height = w, h
    shape.name = name
    sppr = shape._element.spPr
    geometry = sppr.find(f"{{{A}}}custGeom")
    for item in list(geometry):
        geometry.remove(item)
    for tag in ("avLst", "gdLst", "ahLst", "cxnLst"):
        element(geometry, tag)
    element(geometry, "rect", l="0", t="0", r="r", b="b")
    paths = element(geometry, "pathLst")
    precision = 1000
    path = element(paths, "path", w=str(round(parsed.width * precision)),
                   h=str(round(parsed.height * precision)), stroke="0")

    def point(parent, x, y):
        element(parent, "pt", x=str(round(x * precision)),
                y=str(round((parsed.height - parsed.depth - y) * precision)))

    for font, size, code, ox, oy in parsed.glyphs:
        font.set_size(size, 72)
        font.load_char(code)
        vertices, codes = font.get_path()
        current = None
        index = 0
        while index < len(codes):
            code = codes[index]
            vertex = vertices[index] + (ox, oy)
            if code == MathPath.MOVETO:
                point(element(path, "moveTo"), *vertex)
                current = vertex
            elif code == MathPath.LINETO:
                point(element(path, "lnTo"), *vertex)
                current = vertex
            elif code == MathPath.CURVE3:
                end = vertices[index + 1] + (ox, oy)
                cubic = element(path, "cubicBezTo")
                point(cubic, *(current + (vertex - current) * 2 / 3))
                point(cubic, *(end + (vertex - end) * 2 / 3))
                point(cubic, *end)
                current = end
                index += 1
            elif code == MathPath.CURVE4:
                cubic = element(path, "cubicBezTo")
                for offset in range(3):
                    point(cubic, *(vertices[index + offset] + (ox, oy)))
                current = vertices[index + 2] + (ox, oy)
                index += 2
            elif code == MathPath.CLOSEPOLY:
                element(path, "close")
            index += 1
    for x, y, rw, rh in parsed.rects:
        for tag, coordinates in (("moveTo", (x, y)), ("lnTo", (x + rw, y)),
                                 ("lnTo", (x + rw, y + rh)), ("lnTo", (x, y + rh))):
            point(element(path, tag), *coordinates)
        element(path, "close")
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("spec", type=Path, help="JSON list: page, expression, box_inches, optional color/name")
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    prs = Presentation(args.pptx)
    specs = json.loads(args.spec.read_text(encoding="utf-8"))
    for item in specs:
        page = item["page"]
        if not 1 <= page <= len(prs.slides):
            raise ValueError("formula page is outside deck")
        box = [Inches(value) for value in item["box_inches"]]
        if len(box) != 4 or box[0] < 0 or box[1] < 0 or box[0] + box[2] > prs.slide_width or box[1] + box[3] > prs.slide_height:
            raise ValueError("formula box must fit inside slide")
        add_formula(prs.slides[page - 1], item["expression"], *box,
                    item.get("color", "202020"), item.get("name", "Formula"))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=args.output.parent, prefix=".formula-") as directory:
        temporary = Path(directory) / "deck.pptx"
        prs.save(temporary)
        Presentation(temporary)
        temporary.replace(args.output)
    print(json.dumps({"formulas_added": len(specs), "native_vectors": True}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
