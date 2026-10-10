#!/usr/bin/env python3
"""Copy editable slides between presentations while retaining media and notes."""
from __future__ import annotations

import argparse
import copy
import json
import tempfile
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.opc.packuri import PackURI

R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _remap_relationship_ids(element, mapping: dict[str, str]) -> None:
    """Replace relationship ids in a copied slide XML tree."""
    for node in element.iter():
        for attribute, value in list(node.attrib.items()):
            if attribute.startswith(f"{{{R}}}") and value in mapping:
                node.set(attribute, mapping[value])


def copy_slide(source_slide, target_prs: Presentation):
    """Append one source slide to ``target_prs`` using the target's layout/theme."""
    layout = target_prs.slide_layouts.get_by_name(source_slide.slide_layout.name)
    if layout is None:
        layout = target_prs.slide_layouts[-1]
    target = target_prs.slides.add_slide(layout)
    target_rels = target.part.rels
    relationship_ids: dict[str, str] = {}
    for relationship in source_slide.part.rels.values():
        if relationship.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        relationship_ids[relationship.rId] = target_rels._add_relationship(
            relationship.reltype, relationship._target, relationship.is_external
        )

    copied = copy.deepcopy(source_slide._element)
    _remap_relationship_ids(copied, relationship_ids)
    destination = target._element
    for child in list(destination):
        destination.remove(child)
    for child in copied:
        destination.append(child)

    if source_slide.has_notes_slide and source_slide.notes_slide.notes_text_frame.text:
        target.notes_slide.notes_text_frame.text = source_slide.notes_slide.notes_text_frame.text
    return target


def import_dependencies(source_prs, target_prs, pages):
    """Rename imported parts to avoid ZIP collisions; reject internal slide links.

    Images, charts and their embedded workbooks may share filenames with the
    target. Each imported graph retains sharing but receives a unique partname.
    Slide navigation needs an explicit page mapping and is therefore rejected.
    """
    occupied = {str(part.partname) for part in target_prs.part.package.iter_parts()}
    visited = set()

    def visit(part):
        if part in visited:
            return
        visited.add(part)
        original = str(part.partname)
        name = original
        number = 1
        path = Path(original)
        while name in occupied:
            name = str(path.with_name(f"{path.stem}_copy{number}{path.suffix}"))
            number += 1
        occupied.add(name)
        part._partname = PackURI(name)
        for relationship in part.rels.values():
            if relationship.is_external:
                continue
            if relationship.reltype.endswith("/slide"):
                raise ValueError("internal slide links require an explicit page mapping")
            visit(relationship.target_part)

    for page in pages:
        for relationship in source_prs.slides[page - 1].part.rels.values():
            if relationship.is_external or relationship.reltype.endswith(("/slideLayout", "/notesSlide")):
                continue
            if relationship.reltype.endswith("/slide"):
                raise ValueError("internal slide links require an explicit page mapping")
            visit(relationship.target_part)


def copy_slides(source: Path, target: Path, output: Path, pages: list[int], position: int | None = None) -> dict:
    """Copy selected 1-based pages from ``source`` into ``target`` atomically.

    The target's slide layout and theme remain authoritative. Source media,
    charts, external hyperlinks and editable shape XML are copied with relationship ids
    remapped; source notes are copied when present. ``position`` is a zero-based
    insertion point in the target's existing slide list, defaulting to append.
    """
    if output.resolve() in {source.resolve(), target.resolve()}:
        raise ValueError("output must be separate from both input presentations")
    if not pages or any(type(page) is not int or page < 1 for page in pages):
        raise ValueError("pages must contain positive 1-based integers")
    source_prs = Presentation(source)
    target_prs = Presentation(target)
    if any(page > len(source_prs.slides) for page in pages):
        raise ValueError("source page is outside the presentation")
    if position is not None and not 0 <= position <= len(target_prs.slides):
        raise ValueError("position must be within the target slide list")
    if (source_prs.slide_width, source_prs.slide_height) != (target_prs.slide_width, target_prs.slide_height):
        raise ValueError("source and target slide dimensions must match")

    import_dependencies(source_prs, target_prs, pages)
    copied = [copy_slide(source_prs.slides[page - 1], target_prs) for page in pages]

    def slide_id(slide):
        return next(item for item in target_prs.slides._sldIdLst
                    if target_prs.part.related_part(item.rId).partname == slide.part.partname)

    if position is not None:
        ids = target_prs.slides._sldIdLst
        copied_ids = [slide_id(slide) for slide in copied]
        for copied_id in copied_ids:
            ids.remove(copied_id)
        anchor = list(ids)[position] if position < len(ids) else None
        for copied_id in copied_ids:
            if anchor is None:
                ids.append(copied_id)
            else:
                ids.insert(ids.index(anchor), copied_id)

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix=".slide-ops-") as directory:
        temporary = Path(directory) / "presentation.pptx"
        target_prs.save(temporary)
        with ZipFile(temporary) as archive:
            if len(archive.namelist()) != len(set(archive.namelist())) or archive.testzip():
                raise ValueError("copied deck contains invalid or duplicate ZIP entries")
        readback = Presentation(temporary)
        if len(readback.slides) != len(target_prs.slides):
            raise ValueError("slide count changed during save")
        temporary.replace(output)
    return {"source_pages": pages, "copied": len(pages), "output_pages": len(target_prs.slides)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="presentation containing reusable pages")
    parser.add_argument("target", type=Path, help="presentation receiving the pages")
    parser.add_argument("output", type=Path)
    parser.add_argument("--pages", type=int, nargs="+", required=True)
    parser.add_argument("--position", type=int, help="zero-based insertion point; default appends")
    args = parser.parse_args()
    print(json.dumps(copy_slides(args.source, args.target, args.output, args.pages, args.position)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
