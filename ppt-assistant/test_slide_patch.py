"""Verify targeted layout repairs preserve package content and rich text."""
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from slide_patch import inspect, patch


class SlidePatchTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.pptx"
        self.spec = self.root / "patch.json"
        self.output = self.root / "patched.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text, run.font.bold, run.font.size = "保留格式", True, Pt(18)
        self.shape_id = shape.shape_id
        slide.notes_slide.notes_text_frame.text = "Keep notes"
        image = self.root / "image.png"
        Image.new("RGB", (20, 20), "green").save(image)
        slide.shapes.add_picture(str(image), Inches(1), Inches(3))
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
            Inches(1), Inches(1), Inches(4), Inches(1)).text = "Other page"
        # Reorder to prove logical page numbering, not slideN.xml naming.
        ids = prs.slides._sldIdLst
        ids.insert(0, ids[-1])
        prs.save(self.source)

    def test_selected_xml_only_and_rich_text_retained(self):
        self.assertEqual(inspect(self.source)["pages"][1]["shapes"][0]["shape_id"], self.shape_id)
        self.spec.write_text(json.dumps([{"page": 2, "shape_id": self.shape_id,
            "box_inches": [2, 1, 4, 1.5], "text_frame": {"word_wrap": True, "line_spacing": 1.2,
            "vertical_anchor": "middle", "margin_inches": [0.1, 0.1, 0.1, 0.1],
            "space_before_pt": 0, "space_after_pt": 3, "alignment": "center"}}]), encoding="utf-8")
        report = patch(self.source, self.spec, self.output)
        self.assertTrue(report["other_parts_preserved"])
        with ZipFile(self.source) as before, ZipFile(self.output) as after:
            self.assertEqual(before.namelist(), after.namelist())
            changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
            self.assertEqual(changed, ["ppt/slides/slide1.xml"])
        shape = Presentation(self.output).slides[1].shapes[0]
        self.assertEqual(shape.left, Inches(2))
        self.assertEqual(shape.text, "保留格式")
        self.assertEqual(shape.text_frame.vertical_anchor, MSO_ANCHOR.MIDDLE)
        paragraph = shape.text_frame.paragraphs[0]
        self.assertAlmostEqual(paragraph.line_spacing, 1.2)
        self.assertEqual(paragraph.space_after, Pt(3))
        self.assertTrue(paragraph.runs[0].font.bold)
        self.assertEqual(paragraph.runs[0].font.size, Pt(18))

    def test_invalid_specs_preserve_existing_output(self):
        self.output.write_bytes(b"keep")
        for changes in ({"box_inches": [9, 1, 4, 1]}, {"text_frame": {"line_spacing": 0}},
                        {"text_frame": {"word_wrap": "yes"}}, {"shape_id": 999, "box_inches": [1, 1, 2, 1]},
                        {"box_inches": [1, 1, 2, float("nan")]}, {"text_frame": {"unsupported": True}}):
            with self.subTest(changes=changes):
                operation = {"page": 2, "shape_id": self.shape_id, **changes}
                self.spec.write_text(json.dumps([operation]), encoding="utf-8")
                with self.assertRaises(ValueError):
                    patch(self.source, self.spec, self.output)
                self.assertEqual(self.output.read_bytes(), b"keep")


if __name__ == "__main__":
    unittest.main()
