#!/usr/bin/env python3
"""Convert the complete video to a GIF, or replace a named PPT picture with it."""
from __future__ import annotations

import argparse
import json
import math
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt


def choose_speed(duration: float, target: float = 10, multiple: int = 5) -> int:
    """Pick an integer speed with duration closest to target, preferring multiples."""
    if not math.isfinite(duration) or duration <= 0 or not math.isfinite(target) or target <= 0 or multiple < 1:
        raise ValueError("positive finite duration, target and speed multiple required")
    ideal = duration / target
    if ideal < multiple:
        candidates = range(1, multiple + 1)
    else:
        lower = max(multiple, math.floor(ideal / multiple) * multiple)
        candidates = (lower, lower + multiple)
    return min(candidates, key=lambda speed: (abs(duration / speed - target), speed))


def gif_info(data: bytes) -> dict:
    """Read frame count and actual encoded duration; reject static GIFs."""
    with Image.open(BytesIO(data)) as image:
        if image.format != "GIF" or image.n_frames < 2:
            raise ValueError("expected an animated, multi-frame GIF")
        duration = 0
        for frame in range(image.n_frames):
            image.seek(frame)
            duration += image.info.get("duration", 0)
        return {"frames": image.n_frames, "duration_seconds": duration / 1000, "size": list(image.size)}


def convert(source: Path, output: Path, *, speed: int | None = None,
            target: float = 10, multiple: int = 5, fps: int = 20, size: int = 640,
            timeout: float = 300) -> dict:
    """Decode the entire first video stream; use speed rather than trimming."""
    if source.resolve() == output.resolve() or output.suffix.lower() != ".gif":
        raise ValueError("output must be a separate .gif file")
    if fps < 1 or fps > 50 or size < 2 or timeout <= 0:
        raise ValueError("fps must be 1–50; positive size and timeout required")
    probe = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries",
                            "stream=duration:format=duration", "-of", "json", str(source.resolve())],
                           capture_output=True, text=True, check=True, timeout=timeout)
    metadata = json.loads(probe.stdout)
    if not metadata.get("streams"):
        raise ValueError("input has no video stream")
    raw_duration = metadata["streams"][0].get("duration")
    if raw_duration in (None, "N/A", ""):
        raw_duration = metadata.get("format", {}).get("duration")
    if raw_duration in (None, "N/A", ""):
        raise ValueError("input video duration is unavailable")
    duration = float(raw_duration)
    if not math.isfinite(duration) or duration <= 0:
        raise ValueError("input video duration must be positive and finite")
    selected = choose_speed(duration, target, multiple) if speed is None else speed
    if isinstance(selected, bool) or not isinstance(selected, int) or selected < 1:
        raise ValueError("speed must be a positive integer")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix=".gif-") as directory:
        temporary = Path(directory) / "encoded.gif"
        filters = (f"[0:v:0]setpts=(PTS-STARTPTS)/{selected},fps={fps},"
                   f"scale={size}:{size}:force_original_aspect_ratio=decrease:flags=lanczos,"
                   "split[a][b];[a]palettegen=stats_mode=full[p];[b][p]paletteuse=dither=sierra2_4a")
        subprocess.run(["ffmpeg", "-v", "error", "-nostdin", "-i", str(source.resolve()),
                        "-filter_complex", filters, "-loop", "0", str(temporary)],
                       check=True, timeout=timeout)
        info = gif_info(temporary.read_bytes())
        temporary.replace(output)
    return {"source_duration_seconds": duration, "speed": selected, "full_source": True, **info}


def insert(pptx: Path, gif: Path, output: Path, page: int, name: str, speed: int,
           label_size: float = 24, color: str = "FFFFFF", background: str = "C4262E") -> dict:
    """Replace one picture in place, retain aspect ratio, add editable nX label."""
    if speed < 1 or label_size <= 0:
        raise ValueError("positive speed and label size required")
    info = gif_info(gif.read_bytes())
    prs = Presentation(pptx)
    if not 1 <= page <= len(prs.slides):
        raise ValueError("page is outside deck")
    slide = prs.slides[page - 1]
    targets = [shape for shape in slide.shapes if shape.name == name]
    if len(targets) != 1 or targets[0].shape_type != 13:
        raise ValueError("expected exactly one named picture")
    old = targets[0]
    width, height = info["size"]
    scale = min(old.width / width, old.height / height)
    w, h = int(width * scale), int(height * scale)
    x, y = old.left + (old.width - w) // 2, old.top + (old.height - h) // 2
    picture = slide.shapes.add_picture(str(gif), x, y, w, h)
    picture.name = name
    old._element.addprevious(picture._element)
    old._element.getparent().remove(old._element)
    marker = f"{name} speed"
    for shape in list(slide.shapes):
        if shape.name == marker:
            shape._element.getparent().remove(shape._element)
    label = slide.shapes.add_textbox(x, y, Pt(max(60, len(str(speed)) * label_size + 35)), Pt(label_size * 1.6))
    label.name = marker
    label.fill.solid()
    label.fill.fore_color.rgb = RGBColor.from_string(background)
    label.text_frame.text = f"{speed}X"
    paragraph = label.text_frame.paragraphs[0]
    paragraph.font.size = Pt(label_size)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor.from_string(color)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix=".gif-deck-") as directory:
        temporary = Path(directory) / "deck.pptx"
        prs.save(temporary)
        readback = Presentation(temporary).slides[page - 1]
        actual = next(s for s in readback.shapes if s.name == name)
        if actual.image.blob != gif.read_bytes():
            raise ValueError("GIF did not survive PPT round trip")
        temporary.replace(output)
    return {"page": page, "shape": name, "speed": speed, **info}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    encoder = commands.add_parser("convert")
    encoder.add_argument("source", type=Path)
    encoder.add_argument("output", type=Path)
    encoder.add_argument("--speed", type=int)
    encoder.add_argument("--target-seconds", type=float, default=10)
    encoder.add_argument("--speed-multiple", type=int, default=5)
    encoder.add_argument("--fps", type=int, default=20)
    encoder.add_argument("--size", type=int, default=640)
    encoder.add_argument("--timeout", type=float, default=300)
    inserter = commands.add_parser("insert")
    inserter.add_argument("pptx", type=Path)
    inserter.add_argument("gif", type=Path)
    inserter.add_argument("output", type=Path)
    inserter.add_argument("--page", type=int, required=True)
    inserter.add_argument("--shape", required=True)
    inserter.add_argument("--speed", type=int, required=True)
    inserter.add_argument("--label-size", type=float, default=24)
    inserter.add_argument("--color", default="FFFFFF")
    inserter.add_argument("--background", default="C4262E")
    args = parser.parse_args()
    if args.command == "convert":
        report = convert(args.source, args.output, speed=args.speed, target=args.target_seconds,
                         multiple=args.speed_multiple, fps=args.fps, size=args.size, timeout=args.timeout)
    else:
        report = insert(args.pptx, args.gif, args.output, args.page, args.shape, args.speed,
                        args.label_size, args.color, args.background)
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
