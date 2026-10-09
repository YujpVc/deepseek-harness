"""Keyless CLI checks for template-preserving, editable review decks."""
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from PIL import Image
import numpy as np
from pptx import Presentation
from pptx.util import Inches


class ProbationDeckTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.template = self.base / "template.pptx"
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
        cover = prs.slides.add_slide(prs.slide_layouts[0])
        cover.shapes.title.text = "Original cover"
        prs.slides.add_slide(prs.slide_layouts[1])
        prs.save(self.template)
        Image.new("RGB", (80, 80), "red").save(self.base / "source.png")
        self.spec = {
            "meta": {"title": "Review", "theme": "paper", "deck_mode": "native_v2"},
            "slides": [{"type": "cover"}, {
                "layout": "evidence", "title": "Results", "image": "source.png",
                "charts": [{"title": "Accepted", "before": [17, 26, 39],
                            "after": [45, 59, 51]},
                           {"title": "Narrow", "before": [20, 26, 48],
                            "after": [3, 5, 1]}],
                "note": "Saved frames only", "notes": "Not a success rate",
            }],
        }
        self.output = self.base / "result.pptx"

    def run_cli(self):
        spec_path = self.base / "spec.json"
        spec_path.write_text(json.dumps(self.spec), encoding="utf-8")
        return subprocess.run([
            sys.executable, str(Path(__file__).with_name("build_deck.py")),
            str(spec_path), str(self.output), "--template", str(self.template), "--probation",
        ], capture_output=True, text=True, check=False)

    def test_cli_retains_cover_and_editable_source_values(self):
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        source, output = Presentation(self.template), Presentation(self.output)
        self.assertEqual(source.slides[0]._element.xml, output.slides[0]._element.xml)
        self.assertEqual(len(output.slides), 2)
        charts = [shape.chart for shape in output.slides[1].shapes if shape.has_chart]
        self.assertEqual(len(charts), 2)
        self.assertEqual(list(charts[0].series[1].values), [45, 59, 51])
        self.assertEqual(list(charts[1].series[1].values), [3, 5, 1])
        self.assertEqual(output.slides[1].notes_slide.notes_text_frame.text, "Not a success rate")

    def test_duplicate_photo_copy_is_rejected_without_overwriting_output(self):
        (self.base / "copy.png").write_bytes((self.base / "source.png").read_bytes())
        self.spec["slides"].append({
            "layout": "field", "title": "Field",
            "images": [{"path": "copy.png", "label": "Copy", "caption": ""}], "flow": "",
        })
        self.output.write_bytes(b"existing deck")
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("reuses image", result.stderr)
        self.assertEqual(self.output.read_bytes(), b"existing deck")

    def test_unknown_layout_is_rejected(self):
        self.spec["slides"][1]["layout"] = "typo"
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("unsupported probation layout", result.stderr)
        self.assertFalse(self.output.exists())

    def test_crop_reads_original_photo_without_workspace_intermediates(self):
        self.spec["slides"][1]["image_crop"] = [10, 20, 60, 70]
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        photo = next(s for s in Presentation(self.output).slides[1].shapes if s.shape_type == 13)
        self.assertEqual(photo.image.size, (50, 50))

    def test_out_of_bounds_crop_is_rejected(self):
        self.spec["slides"][1]["image_crop"] = [0, 0, 100, 100]
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("invalid pixel crop", result.stderr)
        self.assertFalse(self.output.exists())

    def test_incompatible_template_size_is_rejected(self):
        template = Presentation(self.template)
        template.slide_width = Inches(10)
        template.save(self.template)
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("13.333", result.stderr)
        self.assertFalse(self.output.exists())

    def test_paired_grasps_keep_six_photos_and_editable_wireframes(self):
        camera = {"color_k": [[100, 0, 40], [0, 100, 40], [0, 0, 1]]}
        (self.base / "camera.json").write_text(json.dumps(camera))
        transform = [[1, 0, 0, 0], [0, 1, 0, 0], [0, 0, 1, 1], [0, 0, 0, 1]]
        for phase, width in [("before", .02), ("after", .06)]:
            (self.base / f"{phase}.json").write_text(json.dumps({
                "T_camera_cgn_gripper": transform, "gripper_opening_m": width}))
        frames = []
        for i in range(3):
            filename = f"round{i}.png"
            Image.new("RGB", (80, 80), (90 + i, 100, 100)).save(self.base / filename)
            frames.append({"label": f"Round {i+1}", "image": filename,
                           "calibration": "camera.json", "crop": [0, 0, 80, 80],
                           "before": "before.json", "after": "after.json",
                           "before_opening_mm": 20, "after_opening_mm": 60})
        self.spec["slides"][1] = {"layout": "grasp_comparison", "title": "Paired poses",
                                    "frames": frames, "note": "Recorded poses"}
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        slide = Presentation(self.output).slides[1]
        self.assertEqual(sum(s.shape_type == 13 for s in slide.shapes), 6)
        self.assertEqual(sum(s.shape_type == 9 for s in slide.shapes), 162)
        # A mismatched statistic must fail before replacing the existing deck.
        existing = self.output.read_bytes()
        frames[0]["after_opening_mm"] = 80
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("does not match comparison statistics", result.stderr)
        self.assertEqual(self.output.read_bytes(), existing)

    def test_gallery_keeps_aligned_transparent_source_layers(self):
        frames = []
        for i in range(3):
            photo, layer = f"capture{i}.png", f"layer{i}.png"
            Image.new("RGB", (80, 60), (80 + i, 90, 100)).save(self.base / photo)
            overlay = Image.new("RGBA", (80, 60), (0, 0, 0, 0))
            overlay.putpixel((30, 20), (255, 0, 0, 210))
            overlay.save(self.base / layer)
            frames.append({"title": f"Capture {i}", "body": "Recorded candidates",
                           "image": photo, "layers": [layer], "crop": [10, 10, 70, 50]})
        self.spec["slides"][1] = {"layout": "visual_gallery", "title": "Gallery",
                                    "items": frames, "note": "Saved captures"}
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        pictures = [s for s in Presentation(self.output).slides[1].shapes if s.shape_type == 13]
        self.assertEqual(len(pictures), 6)
        for photo, layer in zip(pictures[::2], pictures[1::2]):
            self.assertEqual((photo.left, photo.top, photo.width, photo.height),
                             (layer.left, layer.top, layer.width, layer.height))
            self.assertEqual(layer.image.size, (60, 40))
            with Image.open(io.BytesIO(layer.image.blob)) as overlay:
                self.assertEqual(overlay.getpixel((20, 10)), (255, 0, 0, 210))
                self.assertEqual(overlay.getpixel((0, 0))[3], 0)
        existing = self.output.read_bytes()
        Image.new("RGBA", (79, 60), (255, 0, 0, 100)).save(self.base / "layer0.png")
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("matching the source photo dimensions", result.stderr)
        self.assertEqual(self.output.read_bytes(), existing)

    def test_depth_case_pairs_candidates_and_depth_in_one_page(self):
        transform = [[1, 0, 0, 0], [0, 1, 0, 0], [0, 0, 1, 1], [0, 0, 0, 1]]
        (self.base / "camera.json").write_text(json.dumps({
            "color_k": [[100, 0, 40], [0, 100, 40], [0, 0, 1]]}))
        np.save(self.base / "before.npy", np.ones((80, 80)))
        np.save(self.base / "after.npy", np.ones((80, 80)))
        for phase, count in (("before", 3), ("after", 4)):
            (self.base / f"{phase}.json").write_text(json.dumps({"grasps": {
                "total_grasps": count, "pred_grasps_cam": {"0": [transform] * count},
                "scores": {"0": [i / 10 for i in range(count)]},
                "gripper_openings": {"0": [.04] * count}}}))
        self.spec["slides"][1] = {"layout": "depth_case", "title": "Depth optimization",
            "image": "source.png", "crop": [0, 0, 80, 80], "calibration": "camera.json",
            "before_response": "before.json", "after_response": "after.json",
            "before_array": "before.npy", "after_array": "after.npy",
            "depth_crop": [0, 0, 80, 80], "limits_m": [.9, 1.1],
            "before_count": 3, "after_count": 4, "before_points": 10, "after_points": 12,
            "solution": "Complete depth\nRecompute poses", "note": "Same source"}
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        slide = Presentation(self.output).slides[1]
        self.assertFalse(any(s.is_placeholder for s in slide.shapes))
        self.assertEqual(sum(s.shape_type == 13 for s in slide.shapes), 4)
        lines = [s for s in slide.shapes if s.shape_type == 9]
        self.assertEqual(len(lines), 162)
        self.assertEqual(len({(str(s.line.color.rgb), s.line.width) for s in lines}), 1)
        existing = self.output.read_bytes()
        self.spec["slides"][1]["after_count"] = 7
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("candidate count does not match source", result.stderr)
        self.assertEqual(self.output.read_bytes(), existing)


if __name__ == "__main__":
    unittest.main()
